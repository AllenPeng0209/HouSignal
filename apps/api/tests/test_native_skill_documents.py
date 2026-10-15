"""Default-suite real rendering fixtures for native document and visual tools."""
from pathlib import Path
import json
import sys
import xml.etree.ElementTree as ET

import pytest
from PIL import Image, ImageChops, ImageDraw

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

try:
    from app.content_skills.tools import documents
except ImportError:
    documents = None


def run(tmp_path, name, text='', parameters=None, assets=()):
    assert documents is not None, 'native document dispatch has not been implemented'
    return documents.dispatch(name, input_text=text, assets=assets,
        parameters=parameters or {}, output_dir=tmp_path / 'out', context={})


def artifacts(result, suffix):
    return [p for p in result['artifact_paths'] if p.suffix == suffix]


@pytest.mark.parametrize('name,size', [
    ('card-design', (320, 480)), ('card-xiaohongshu', (320, 480)),
    ('card-quote', (480, 270)), ('poster-hero', (320, 568)),
])
def test_cards_render_content_in_real_png_and_matching_local_svg(tmp_path, name, size):
    result = run(tmp_path, name, '# Native visual\n\nA real point\n- Keep the source\n- Read the result',
        {'width': size[0], 'height': size[1], 'scale': 1, 'style': 'swiss'})
    with Image.open(artifacts(result, '.png')[0]) as img:
        assert img.size == size
        assert ImageChops.difference(img.convert('RGB'), Image.new('RGB', size, img.getpixel((0, 0)))).getbbox()
    svg = artifacts(result, '.svg')[0].read_text()
    root = ET.fromstring(svg)
    assert 'Native visual' in ''.join(root.itertext())
    html = artifacts(result, '.html')[0].read_text()
    assert '<svg' in html and 'cdn.' not in html


def test_knowledge_cards_do_not_drop_pages(tmp_path):
    result = run(tmp_path, 'card-xiaohongshu', '', {'cards': [
        {'title': 'First', 'body': ['One']}, {'title': 'Second', 'body': ['Two']}],
        'width': 320, 'height': 480, 'scale': 1})
    assert [p.name for p in artifacts(result, '.png')] == ['card_1.png', 'card_2.png']
    assert 'Second' in artifacts(result, '.svg')[1].read_text()


def test_explicit_card_title_preserves_a_plain_first_body_paragraph(tmp_path):
    body = '清晰的标题、克制的文案、可核验的事实。'
    result = run(tmp_path, 'poster-hero', body + '\n原创示意内容，用于流程验收。', {
        'title': '看得清，才有下一步', 'width': 1080, 'height': 1920, 'scale': 1})
    root = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    rendered = ''.join(root.itertext())
    assert body in rendered
    assert '看得清，才有下一步' in rendered


def test_explicit_card_title_may_replace_a_markdown_heading_without_dropping_body(tmp_path):
    result = run(tmp_path, 'card-design', '# 旧标题\n\n正文事实必须保留。', {
        'title': '已确认的新标题', 'width': 1080, 'height': 1440, 'scale': 1})
    rendered = ''.join(ET.fromstring(artifacts(result, '.svg')[0].read_text()).itertext())
    assert '旧标题' not in rendered
    assert '已确认的新标题' in rendered and '正文事实必须保留。' in rendered


@pytest.mark.parametrize('text,character_width', [
    ('清晰的标题、克制的文案、可核验的事实。', 9.55),
    ('内容（必须保留）和证据。', 3.05),
    ('章节「保留证据」再核对。', 3.05),
    ('读者“需要明确信息”，不能编造。', 3.05),
    ('范围【数据（确认）】与引用。', 4.05),
])
def test_chinese_soft_wrap_preserves_characters_and_forbids_hanging_punctuation(text, character_width):
    size=28; font=documents._font(size)
    width=font.getlength('数')*character_width
    lines=documents._wrap(text,width,size)
    closing=set('，。！？：；、）】》」』”’')
    opening=set('（【《「『“‘')
    assert ''.join(lines)==text
    assert all(line and line[0] not in closing and line[-1] not in opening for line in lines), lines
    assert all(font.getlength(line)<=width for line in lines)


def test_chinese_wrap_rejects_impossible_box_instead_of_overflow_or_losing_a_character():
    font=documents._font(28)
    with pytest.raises(ValueError,match='画幅|换行|宽度'):
        documents._wrap('数。',font.getlength('数')+.1,28)


def test_actual_poster_copy_has_no_punctuation_only_line_and_stays_inside_body_box(tmp_path):
    body='清晰的标题、克制的文案、可核验的事实。'
    result=run(tmp_path,'poster-hero',body+'\n原创示意内容，用于流程验收。',{
        'title':'看得清，才有下一步','subtitle':'原创杯具示意 · 不代表实际商品参数',
        'cta':'查看项目资料','width':1080,'height':1920,'scale':1,'style':'swiss'})
    root=ET.fromstring(artifacts(result,'.svg')[0].read_text())
    nodes=list(root.iter('{http://www.w3.org/2000/svg}text'))
    assert not any((node.text or '').startswith(tuple('，。！？：；、）】》」』”’')) for node in nodes)
    assert body in ''.join(root.itertext())
    for node in nodes:
        size=int(float(node.attrib['font-size']));font=documents._font(size)
        assert float(node.attrib['x'])+font.getlength(node.text or '')<=1020.1


def test_comparison_renders_given_values_without_inventing_winner(tmp_path):
    result = run(tmp_path, 'comparison-card', '', {'data': [
        {'Dimension': 'Price', 'A': 100, 'B': 120},
        {'Dimension': 'Storage', 'A': '64GB', 'B': '128GB'}],
        'width': 480, 'height': 640, 'scale': 1})
    root = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    text = ''.join(root.itertext())
    assert all(value in text for value in ['Price', '100', '120', '64GB', '128GB'])
    assert 'winner' not in text


def test_mindmap_preserves_heading_and_list_hierarchy_offline(tmp_path):
    result = run(tmp_path, 'mindmap', '# Plan\n## Research\n- Source A\n- Source B\n## Produce\n- Render',
        {'width': 700, 'height': 400, 'scale': 1})
    tree = json.loads(artifacts(result, '.json')[0].read_text())
    assert tree['name'] == 'Plan'
    assert [child['name'] for child in tree['children']] == ['Research', 'Produce']
    assert tree['children'][0]['children'][0]['name'] == 'Source A'
    html = artifacts(result, '.html')[0].read_text()
    assert 'zoom' in html and 'src="https:' not in html


def test_markdown_convert_renders_tables_code_and_escaped_html(tmp_path):
    result = run(tmp_path, 'doc-convert', '# Report\n\n**Important** and `code`\n\n| City | Value |\n|---|---|\n| A | 7 |\n\n```python\nprint(7)\n```\n\n<script>alert(1)</script>',
        {'format': 'html'})
    html = artifacts(result, '.html')[0].read_text()
    assert '<strong>Important</strong>' in html and '<table>' in html
    assert '<pre' in html and '<code>print(7)</code>' in html and '&lt;script&gt;alert' in html
    assert '<script>alert' not in html


def test_markdown_png_has_complete_long_document(tmp_path):
    result = run(tmp_path, 'doc-convert', '# Long\n\n' + '\n\n'.join(f'Paragraph {n}' for n in range(25)),
        {'format': 'png', 'width': 400, 'scale': 1})
    with Image.open(artifacts(result, '.png')[0]) as img:
        assert img.width == 400 and img.height > 800
    assert 'Paragraph 24' in artifacts(result, '.html')[0].read_text()


def test_markdown_pdf_is_real_and_contains_document_text(tmp_path):
    pytest.importorskip('reportlab')
    from pypdf import PdfReader
    result = run(tmp_path, 'doc-convert', '# Document\n\nA real paragraph.\n\n- Point one\n- Point two', {'format': 'pdf'})
    path = artifacts(result, '.pdf')[0]
    assert path.read_bytes().startswith(b'%PDF-')
    text = '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)
    assert 'Document' in text and 'Point two' in text


def test_chinese_pdf_embeds_its_unicode_font_for_portable_real_glyphs(tmp_path):
    from pypdf import PdfReader
    result = run(tmp_path, 'doc-convert', '# 原创内容交付说明\n\n- 保留真实文字与结构。', {'format': 'pdf'})
    reader = PdfReader(artifacts(result, '.pdf')[0])
    assert '保留真实文字与结构' in reader.pages[0].extract_text()
    embedded = []
    for reference in reader.pages[0]['/Resources']['/Font'].values():
        font = reference.get_object()
        descriptors = [font.get('/FontDescriptor')]
        descriptors.extend(child.get_object().get('/FontDescriptor') for child in font.get('/DescendantFonts', []))
        for reference in descriptors:
            if reference:
                descriptor = reference.get_object()
                embedded.extend(key for key in ('/FontFile','/FontFile2','/FontFile3') if key in descriptor)
    assert embedded, 'Portable PDF must include font glyphs, rather than depend on a viewer-specific STSong substitute'


@pytest.mark.parametrize('chart_type', ['bar', 'column', 'line', 'area', 'pie', 'radar', 'funnel'])
def test_charts_render_user_csv_data_to_png_and_svg(tmp_path, chart_type):
    result = run(tmp_path, 'chart-visualization', 'category,value\nAlpha,7\nBeta,13\nGamma,9\n',
        {'chart_type': chart_type, 'width': 480, 'height': 320, 'scale': 1})
    assert len(artifacts(result, '.png')) == 1 and len(artifacts(result, '.svg')) == 1
    data = json.loads(artifacts(result, '.json')[0].read_text())
    assert [row['value'] for row in data['data']] == [7, 13, 9]
    with Image.open(artifacts(result, '.png')[0]) as img:
        assert img.size == (480, 320)


def test_scatter_validates_xy_and_uses_real_points(tmp_path):
    result = run(tmp_path, 'chart-visualization', '', {'chart_type': 'scatter',
        'data': [{'x': 2, 'y': 3}, {'x': 8, 'y': 9}], 'width': 400, 'height': 300, 'scale': 1})
    assert len(artifacts(result, '.svg')) == 1
    assert '<ellipse' in artifacts(result, '.svg')[0].read_text()


def test_horizontal_bar_numeric_scale_is_on_x_axis_and_value_labels_are_not_clipped(tmp_path):
    result = run(tmp_path, 'chart-visualization', '', {'chart_type': 'bar',
        'data': [{'week': '第一周', 'completed': 12}, {'week': '第二周', 'completed': 18},
                 {'week': '第三周', 'completed': 15}], 'x': 'week', 'y': 'completed',
        'width': 1200, 'height': 800, 'scale': 1})
    root = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    texts = list(root.iter('{http://www.w3.org/2000/svg}text'))
    plot_bottom = 710
    scale = [node for node in texts if node.text in {'0', '4.5', '9', '13.5', '18'} and float(node.attrib['y']) >= plot_bottom]
    assert len(scale) == 5, 'Horizontal numeric ticks must follow the x-axis below the bars'
    assert len({node.attrib['x'] for node in scale}) == 5
    for node in texts:
        if node.text == '18' and float(node.attrib['y']) < plot_bottom:
            font = documents._font(int(float(node.attrib['font-size'])), 'sans')
            assert float(node.attrib['x']) + font.getlength(node.text) <= 1188


def test_data_report_from_asset_calculates_observed_statistics(tmp_path):
    src = tmp_path / 'sales.csv'
    src.write_text('region,sales,cost\nEast,10,4\nEast,20,8\nWest,30,12\n')
    result = run(tmp_path, 'data-report', assets=[{'path': src}], parameters={'width': 500, 'scale': 1})
    overview = json.loads(artifacts(result, '.json')[0].read_text())
    assert overview['shape'] == {'rows': 3, 'cols': 3}
    assert overview['numeric_stats']['sales']['sum'] == 60
    assert overview['numeric_stats']['cost']['mean'] == 8
    html = artifacts(result, '.html')[0].read_text()
    assert '<table>' in html and '60' in html and '<svg' in html


def test_infographic_flow_renders_steps_and_arrows(tmp_path):
    result = run(tmp_path, 'infographic', '# Delivery\n- Research\n- Produce\n- Review',
        {'op': 'flow', 'width': 640, 'height': 400, 'scale': 1})
    svg = artifacts(result, '.svg')[0].read_text()
    assert all(word in svg for word in ['Research', 'Produce', 'Review'])
    assert '<polygon' in svg


def test_infographic_animation_preserves_final_target_and_multiple_frames(tmp_path):
    result = run(tmp_path, 'infographic', '', {'op': 'count-up',
        'data': {'items': [{'label': 'Users', 'value': 125}]},
        'width': 320, 'height': 200, 'fps': 4, 'duration': 0.5, 'scale': 1})
    with Image.open(artifacts(result, '.gif')[0]) as img:
        assert img.n_frames >= 2 and img.size == (320, 200)
    assert json.loads(artifacts(result, '.json')[0].read_text())['data']['items'][0]['value'] == 125


def test_card_audit_rejects_blank_image_and_accepts_dense_content(tmp_path):
    blank = tmp_path / 'blank.png'
    dense = tmp_path / 'dense.png'
    Image.new('RGB', (200, 200), 'white').save(blank)
    img = Image.new('RGB', (200, 200), 'white')
    draw = ImageDraw.Draw(img)
    for x in range(8, 195, 6):
        draw.line((x, 8, x, 192), fill='black')
    img.save(dense)
    result = run(tmp_path, 'card-design', parameters={'op': 'audit'}, assets=[blank, dense])
    report = json.loads(artifacts(result, '.json')[0].read_text())
    assert report[0]['pass'] is False and report[1]['pass'] is True
    assert report[1]['span'] > 0.9


@pytest.mark.parametrize('parameters,text', [
    ({'output': '../escape.png'}, 'Some content'),
    ({'width': -1}, 'Some content'),
    ({'scale': 20}, 'Some content'),
    ({'style': 'unrecognized'}, 'Some content'),
    ({}, ''),
])
def test_invalid_card_parameters_write_no_artifacts(tmp_path, parameters, text):
    with pytest.raises(ValueError):
        run(tmp_path, 'card-quote', text, parameters)
    assert not list(tmp_path.rglob('*.png'))


def test_nonfinite_chart_values_are_rejected_before_any_output(tmp_path):
    with pytest.raises(ValueError, match='finite|有限'):
        run(tmp_path, 'chart-visualization', parameters={'data': [{'category': 'A', 'value': float('nan')}]})
    assert not list(tmp_path.rglob('*.png'))


def test_artifact_symlink_cannot_overwrite_unrelated_file(tmp_path):
    outside = tmp_path / 'outside.txt'
    outside.write_text('keep')
    out = tmp_path / 'out'
    out.mkdir()
    (out / 'card.png').symlink_to(outside)
    with pytest.raises(ValueError, match='symbolic|symlink|链接'):
        run(tmp_path, 'card-quote', 'Safe card', {'width': 320, 'height': 180, 'scale': 1})
    assert outside.read_text() == 'keep'


def test_gzh_design_creates_safe_publishable_article_html(tmp_path):
    result = run(tmp_path, 'gzh-design', '# Article\n\nContent with **emphasis**.', {'theme': 'olive-journal'})
    html = artifacts(result, '.html')[0].read_text()
    assert 'Article' in html and '<strong>emphasis</strong>' in html
    assert 'olive-journal' in html and 'cdn.' not in html


def test_document_readiness_is_independent_of_easel_service():
    assert documents is not None, 'native document readiness has not been implemented'
    for name in ['card-design', 'card-quote', 'poster-hero', 'chart-visualization', 'mindmap', 'doc-convert']:
        assert documents.readiness(name)['status'] == 'ready'
    assert documents.readiness('unknown')['status'] == 'unavailable'


@pytest.mark.parametrize('kind,data', [
    ('waterfall', [{'category': 'Start', 'value': 10}, {'category': 'Gain', 'value': 5}, {'category': 'Loss', 'value': -3}, {'category': 'Total', 'isTotal': True}]),
    ('dual-axes', {'categories': ['Jan', 'Feb'], 'series': [{'type': 'column', 'data': [10, 20]}, {'type': 'line', 'data': [1, 3]}]}),
    ('boxplot', [{'category': 'A', 'value': 1}, {'category': 'A', 'value': 2}, {'category': 'A', 'value': 8}, {'category': 'B', 'value': 3}, {'category': 'B', 'value': 9}]),
    ('violin', [{'category': 'A', 'value': 1}, {'category': 'A', 'value': 2}, {'category': 'A', 'value': 8}, {'category': 'B', 'value': 3}, {'category': 'B', 'value': 9}]),
    ('liquid', {'percent': .65}),
    ('word-cloud', [{'text': 'Source', 'value': 10}, {'text': 'Review', 'value': 6}, {'text': 'Render', 'value': 4}]),
    ('treemap', [{'name': 'A', 'value': 3}, {'name': 'B', 'value': 7, 'children': [{'name': 'B1', 'value': 4}, {'name': 'B2', 'value': 3}]}]),
    ('sankey', [{'source': 'Start', 'target': 'A', 'value': 4}, {'source': 'Start', 'target': 'B', 'value': 6}]),
    ('venn', [{'sets': ['A'], 'value': 5}, {'sets': ['B'], 'value': 6}, {'sets': ['A', 'B'], 'value': 2}]),
    ('network-graph', {'nodes': [{'name': 'A'}, {'name': 'B'}], 'edges': [{'source': 'A', 'target': 'B', 'name': 'knows'}]}),
    ('flow-diagram', {'nodes': [{'name': 'A'}, {'name': 'B'}], 'edges': [{'source': 'A', 'target': 'B'}]}),
    ('mind-map', {'name': 'Root', 'children': [{'name': 'One'}, {'name': 'Two'}]}),
    ('organization-chart', {'name': 'Leader', 'children': [{'name': 'Team One'}, {'name': 'Team Two'}]}),
    ('fishbone-diagram', {'name': 'Problem', 'children': [{'name': 'People', 'children': [{'name': 'Training'}]}, {'name': 'Process'}]}),
])
def test_advanced_chart_shapes_render_source_data_contracts(tmp_path, kind, data):
    result = run(tmp_path, 'chart-visualization', parameters={'chart_type': kind, 'data': data,
        'width': 640, 'height': 400, 'scale': 1})
    assert artifacts(result, '.png')[0].is_file()
    root = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    assert len(list(root)) > 2
    assert json.loads(artifacts(result, '.json')[0].read_text())['chart_type'] == kind


def test_grouped_line_chart_does_not_connect_different_series(tmp_path):
    result = run(tmp_path, 'chart-visualization', parameters={'chart_type': 'line', 'data': [
        {'time': 'Jan', 'value': 1, 'group': 'A'}, {'time': 'Feb', 'value': 2, 'group': 'A'},
        {'time': 'Jan', 'value': 7, 'group': 'B'}, {'time': 'Feb', 'value': 6, 'group': 'B'}],
        'width': 500, 'height': 300, 'scale': 1})
    svg = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    colored_lines = [element for element in svg if element.tag.endswith('line') and element.attrib['stroke'] in {'#5B8FF9', '#61DDAA'}]
    assert len(colored_lines) == 2
    assert len({element.attrib['stroke'] for element in colored_lines}) == 2


@pytest.mark.parametrize('parameters', [{'op': []}, {'style': []}, {'format': {}}, {7: True}, {'cards': [{'title': 'A', 'body': [123]}]}])
def test_malformed_parameter_types_raise_validation_error(tmp_path, parameters):
    with pytest.raises(ValueError):
        run(tmp_path, 'card-design', 'Content', parameters)


def test_nested_output_symlink_is_rejected(tmp_path):
    assert documents is not None
    outside = tmp_path / 'outside'
    outside.mkdir()
    root = tmp_path / 'link'
    root.symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match='symbolic|symlink|链接'):
        documents.dispatch('card-quote', input_text='A real quote', assets=[], parameters={},
            output_dir=root / 'nested', context={})
    assert not list(outside.rglob('*'))


def test_docx_output_contains_all_real_markdown_content(tmp_path):
    from docx import Document
    result = run(tmp_path, 'doc-convert', '# Header\n\nBody\n\n| Key | Value |\n|---|---|\n| A | 7 |', {'format': 'docx'})
    doc = Document(artifacts(result, '.docx')[0])
    assert [paragraph.text for paragraph in doc.paragraphs] == ['Header', 'Body']
    assert doc.tables[0].cell(1, 1).text == '7'


def test_all_animation_operations_generate_real_multiframe_gif(tmp_path):
    cases = [
        ('bar-race', {'times': ['2025', '2026'], 'series': {'A': [1, 3], 'B': [2, 4]}}),
        ('line-grow', {'x': ['Jan', 'Feb'], 'series': {'A': [1, 3], 'B': [2, 4]}}),
        ('progress', {'label': 'Complete', 'value': 65, 'max': 100}),
    ]
    for index, (op, data) in enumerate(cases):
        result = run(tmp_path / str(index), 'infographic', parameters={'op': op, 'data': data,
            'width': 320, 'height': 200, 'scale': 1, 'fps': 4, 'duration': .5})
        with Image.open(artifacts(result, '.gif')[0]) as image:
            assert image.n_frames >= 2


def test_org_chart_vertical_orientation_changes_hierarchy_coordinates(tmp_path):
    data = {'name': 'CEO', 'children': [{'name': 'One'}, {'name': 'Two'}]}
    result = run(tmp_path, 'chart-visualization', parameters={'chart_type': 'organization-chart',
        'orient': 'vertical', 'data': data, 'width': 600, 'height': 400, 'scale': 1})
    svg = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    texts = {element.text: element.attrib for element in svg if element.tag.endswith('text')}
    assert float(texts['CEO']['y']) < float(texts['One']['y'])
    assert texts['One']['y'] == texts['Two']['y']


def test_waterfall_uses_cumulative_total_in_geometry(tmp_path):
    result = run(tmp_path, 'chart-visualization', parameters={'chart_type': 'waterfall', 'data': [
        {'category': 'Start', 'value': 10}, {'category': 'Gain', 'value': 5},
        {'category': 'Total', 'isTotal': True}], 'width': 640, 'height': 400, 'scale': 1})
    root = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    bars = [element for element in root if element.tag.endswith('rect') and element.attrib['fill'] in {'#5B8FF9', '#61DDAA'}]
    assert float(bars[-1].attrib['height']) == 315
    assert float(bars[1].attrib['height']) == 105
    assert '15' in ''.join(root.itertext())


def test_sankey_flow_ribbons_keep_input_weight_ratio(tmp_path):
    result = run(tmp_path, 'chart-visualization', parameters={'chart_type': 'sankey', 'data': [
        {'source': 'Start', 'target': 'A', 'value': 4}, {'source': 'Start', 'target': 'B', 'value': 6}],
        'width': 640, 'height': 400, 'scale': 1})
    root = ET.fromstring(artifacts(result, '.svg')[0].read_text())
    ribbons = [element for element in root if element.tag.endswith('polygon')]
    points = [[tuple(map(float, coordinate.split(','))) for coordinate in element.attrib['points'].split()] for element in ribbons]
    heights = [quad[3][1] - quad[0][1] for quad in points]
    assert heights[0] / heights[1] == pytest.approx(2 / 3)


@pytest.mark.parametrize('parameters', [{'data': {'name': 'Root', 'extra': float('nan')}, 'chart_type': 'mind-map'}, {'data': Path('/tmp/forbidden')}, {'shape': [], 'data': {'percent': .5}, 'chart_type': 'liquid'}])
def test_nested_invalid_chart_data_has_clear_validation_error(tmp_path, parameters):
    with pytest.raises(ValueError):
        run(tmp_path, 'chart-visualization', parameters=parameters)
    assert not list(tmp_path.rglob('*.png'))


def test_poster_distributes_given_points_without_dead_bottom_band(tmp_path):
    result = run(tmp_path, 'poster-hero', '# Project launch\n\n- Real assets tell the story\n- Receipts track publication\n- Verified data explains results',
        {'width': 540, 'height': 960, 'scale': 1, 'cta': 'Start your content plan'})
    audited = run(tmp_path / 'audit', 'card-design', parameters={'op': 'audit'}, assets=artifacts(result, '.png'))
    report = json.loads(artifacts(audited, '.json')[0].read_text())[0]
    assert report['bottom_dead'] <= .15 and report['max_gap'] <= .15 and report['span'] >= .8


@pytest.mark.parametrize('name,text,parameters', [
    ('data-report', 'category,value\nA,1', {'format': 'png'}),
    ('infographic', '', {'op': 'count-up', 'format': 'png', 'data': {'label': 'A', 'value': 1}}),
])
def test_unsupported_formats_are_not_silently_substituted(tmp_path, name, text, parameters):
    with pytest.raises(ValueError, match='format|格式'):
        run(tmp_path, name, text, parameters)
    assert not list(tmp_path.rglob('*.png')) and not list(tmp_path.rglob('*.gif'))


def test_json_report_format_returns_only_observed_statistics(tmp_path):
    result = run(tmp_path, 'data-report', 'category,value\nA,7\nB,13', {'format': 'json'})
    assert len(result['artifact_paths']) == 1
    report = json.loads(result['artifact_paths'][0].read_text())
    assert report['numeric_stats']['value']['sum'] == 20

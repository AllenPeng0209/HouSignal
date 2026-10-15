"""Native document and visual operations migrated from Easel (Apache-2.0).

The fixed card sizes, nine design palettes, density audit, document typography,
report statistics, cubic ease-out and GIF data contracts originate in Easel's
render_card.py, mindmap.py, doc_convert.py, card_audit.py, gif_chart.py and
data-report/report.py. This module executes local Python functions; it imports
no Easel runtime, executes no upstream scripts, and contacts no external service.

Pillow and a shared display list render matching PNG/SVG/HTML without Chromium.
The Markdown parser and gzh-design formatter are independent implementations;
no AGPL gzh-design source, CSS or component code is copied into this module.
Its six named themes are functional options with independently written layouts.
Markdown supports headings, paragraphs, lists, quotes, fenced code and tables;
raw HTML is escaped. PDF uses reportlab, DOCX uses python-docx. The supported
operations and parameters are deliberately explicit: unsupported AntV templates
and chart types raise ValueError instead of returning a substitute or fake URL.
"""
from __future__ import annotations

from collections import Counter
from collections.abc import Mapping, Sequence
from functools import lru_cache
from html import escape
from io import BytesIO, StringIO
from pathlib import Path
import csv
import importlib.util
import json
import math
import os
import re
import statistics
import uuid


SKILL_NAMES = frozenset({'card-design', 'card-xiaohongshu', 'card-quote',
    'poster-hero', 'comparison-card', 'chart-visualization', 'infographic',
    'mindmap', 'doc-convert', 'data-report', 'gzh-design'})
CARD_SIZES = {'card-design': (1080, 1440), 'card-xiaohongshu': (1080, 1440),
    'card-quote': (1600, 900), 'poster-hero': (1080, 1920), 'comparison-card': (1080, 1440)}
STYLES = {
    'swiss': ('#fafaf8', '#0a0a0a', '#737373', '#d8d8d4', '#002FA7', 'sans'),
    'editorial': ('#f3f0e8', '#16251b', '#68625a', '#d8d0c4', '#315d93', 'serif'),
    'chinese-ink': ('#f3efe4', '#1a1a1a', '#6b6357', '#b08d57', '#a8322d', 'serif'),
    'cream': ('#FAF6EF', '#7A6A5C', '#9b8a7b', '#EDE4D3', '#A3B0A0', 'sans'),
    'dopamine': ('#FFD500', '#0a0a0a', '#454545', '#0a0a0a', '#002FA7', 'sans'),
    'luxe-noir': ('#0e0d0c', '#ece2cf', '#9a8c75', '#57442b', '#d4a04a', 'serif'),
    'journal': ('#FBF7F0', '#473e35', '#766958', '#E8C4C4', '#B8D8C8', 'sans'),
    'terminal': ('#0d1117', '#c9d1d9', '#6e7681', '#30363d', '#3fb950', 'mono'),
    'botanical': ('#f6f4ee', '#2f3b2c', '#66765e', '#a7bda0', '#7d9b76', 'sans'),
}
_STYLE_NAMES = {'瑞士极简': 'swiss', '杂志编辑': 'editorial', '新中式墨韵': 'chinese-ink',
    '奶油温柔': 'cream', '多巴胺Y2K': 'dopamine', '高奢黑金': 'luxe-noir',
    '手账贴纸': 'journal', '极客终端': 'terminal', '植物清新': 'botanical'}
GZH_THEMES = {'moyu-green': '#059669', 'red-white': '#DC2626',
    'graphite-minimal': '#52525B', 'zen-whitespace': '#4A5D52',
    'moyu-ticket': '#059669', 'olive-journal': '#ed7b2f'}
CHART_TYPES = frozenset({'bar', 'column', 'line', 'area', 'scatter', 'pie',
    'radar', 'funnel', 'histogram', 'spreadsheet', 'boxplot', 'violin', 'waterfall',
    'dual-axes', 'liquid', 'word-cloud', 'treemap', 'sankey', 'venn',
    'network-graph', 'flow-diagram', 'fishbone-diagram', 'organization-chart', 'mind-map'})
_COLORS = ['#5B8FF9', '#61DDAA', '#F6BD16', '#F08BB4', '#7262FD', '#78D3F8', '#9661BC', '#F6903D']
_COMMON = {'op', 'width', 'height', 'scale', 'title'}
_OPTIONS = {
    **{name: _COMMON | {'style', 'subtitle', 'kicker', 'watermark', 'source',
        'cta', 'cards', 'format'} for name in CARD_SIZES},
    'comparison-card': _COMMON | {'style', 'data', 'columns', 'watermark', 'source', 'subtitle', 'format'},
    'chart-visualization': _COMMON | {'data', 'chart_type', 'type', 'x', 'y', 'group',
        'theme', 'axisXTitle', 'axisYTitle', 'innerRadius', 'binNumber', 'format', 'shape', 'orient'},
    'infographic': _COMMON | {'data', 'style', 'fps', 'duration', 'top_n', 'format'},
    'mindmap': _COMMON | {'bg', 'format', 'data'},
    'doc-convert': _COMMON | {'format', 'page_width'},
    'data-report': _COMMON | {'data', 'kpi', 'format'},
    'gzh-design': _COMMON | {'theme', 'format', 'page_width'},
}
_OPS = {**{name: {'render'} for name in CARD_SIZES},
    'card-design': {'render', 'audit'}, 'chart-visualization': {'render'},
    'infographic': {'list', 'flow', 'comparison', 'hierarchy', 'bar-race', 'count-up', 'progress', 'line-grow'},
    'mindmap': {'make'}, 'doc-convert': {'convert'}, 'data-report': {'report', 'analyze'},
    'gzh-design': {'format'}}


def readiness(skill_name):
    if skill_name not in SKILL_NAMES:
        return {'status': 'unavailable', 'reason': '未注册此原生文档技能'}
    if importlib.util.find_spec('PIL') is None:
        return {'status': 'setup', 'reason': '本地渲染需要 Pillow'}
    return {'status': 'ready', 'reason': '楼响本地 Python 渲染可用；PDF 另需 reportlab'}


def _string(value, name, *, maximum=10000, empty=True):
    if not isinstance(value, str) or len(value) > maximum or (not empty and not value.strip()):
        raise ValueError(f'{name} 必须是有效文字（最多 {maximum} 字符）')
    if any(ord(c) < 32 and c not in '\n\r\t' for c in value):
        raise ValueError(f'{name} 包含非法控制字符')
    return value.strip()


def _number(value, name, low=None, high=None, integer=False):
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise ValueError(f'{name} 必须是有限数字')
    try:
        n = float(value)
    except (ValueError, TypeError, OverflowError):
        raise ValueError(f'{name} 必须是有限数字') from None
    if not math.isfinite(n):
        raise ValueError(f'{name} 必须是有限数字 / finite number')
    if (low is not None and n < low) or (high is not None and n > high) or (integer and not n.is_integer()):
        raise ValueError(f'{name} 超出有效范围')
    return int(n) if integer else n


def _color(value, name='color'):
    if not isinstance(value, str) or not re.fullmatch(r'#[0-9A-Fa-f]{6}', value):
        raise ValueError(f'{name} 必须是 #RRGGBB 色值')
    return value


def _validate(skill, parameters):
    if not isinstance(parameters, Mapping):
        raise ValueError('parameters 必须是对象')
    p = dict(parameters)
    if not all(isinstance(key, str) for key in p):
        raise ValueError('参数名称必须是文字')
    unknown = set(p) - _OPTIONS[skill]
    if unknown:
        raise ValueError('不支持的参数: ' + ', '.join(sorted(unknown)))
    default_op = {'mindmap': 'make', 'doc-convert': 'convert', 'data-report': 'report',
        'infographic': 'list', 'gzh-design': 'format'}.get(skill, 'render')
    p['op'] = p.get('op', default_op)
    if not isinstance(p['op'], str) or p['op'] not in _OPS[skill]:
        raise ValueError('不支持的 op: ' + str(p['op']))
    animated = skill == 'infographic' and p['op'] in {'bar-race', 'count-up', 'progress', 'line-grow'}
    default_size = (1600, 1000) if skill == 'mindmap' else (600, 400) if skill == 'chart-visualization' else (900, 900 if p.get('style', 'ring') == 'ring' and p['op'] == 'progress' else 506) if animated else (800, 600)
    dw, dh = CARD_SIZES.get(skill, default_size)
    p['width'] = _number(p.get('width', dw), 'width', 200, 4096, True)
    p['height'] = _number(p.get('height', dh), 'height', 120, 8192, True)
    p['scale'] = _number(p.get('scale', 1 if animated else 2), 'scale', 1, 3, True)
    if p['width'] * p['height'] * p['scale'] ** 2 > 16_000_000:
        raise ValueError('渲染画幅超过 1600 万像素')
    for key in ('title', 'subtitle', 'kicker', 'watermark', 'source', 'cta', 'axisXTitle', 'axisYTitle'):
        if key in p:
            p[key] = _string(p[key], key, maximum=500)
    if skill in CARD_SIZES:
        if 'style' in p and not isinstance(p['style'], str):
            raise ValueError('style 必须是文字')
        style = _STYLE_NAMES.get(p.get('style'), p.get('style', 'luxe-noir' if skill == 'poster-hero' else 'swiss'))
        if style not in STYLES:
            raise ValueError('未知卡片 style')
        p['style'] = style
    if 'format' in p:
        allowed = {'html', 'png', 'pdf', 'docx'} if skill in {'doc-convert', 'gzh-design'} else {'html', 'json'} if skill == 'data-report' else {'gif', 'all'} if animated else {'all', 'html', 'svg', 'png'}
        if not isinstance(p['format'], str) or p['format'] not in allowed:
            raise ValueError('不支持的 format')
    if 'page_width' in p:
        p['page_width'] = _number(p['page_width'], 'page_width', 200, 2000, True)
    if skill == 'gzh-design' and (not isinstance(p.get('theme', 'moyu-green'), str) or p.get('theme', 'moyu-green') not in GZH_THEMES):
        raise ValueError('未知公众号 theme')
    if skill == 'chart-visualization':
        p['chart_type'] = p.get('chart_type', p.get('type', 'column'))
        if not isinstance(p['chart_type'], str) or p['chart_type'] not in CHART_TYPES:
            raise ValueError('原生图表暂不支持此 chart_type: ' + str(p['chart_type']))
        if not isinstance(p.get('theme', 'default'), str) or p.get('theme', 'default') not in {'default', 'academy', 'dark'}:
            raise ValueError('theme 必须是 default/academy/dark')
        for key in ('x', 'y', 'group'):
            if key in p:
                p[key] = _string(p[key], key, maximum=100, empty=False)
        p['innerRadius'] = _number(p.get('innerRadius', 0), 'innerRadius', 0, 0.9)
        p['binNumber'] = _number(p.get('binNumber', 10), 'binNumber', 1, 100, True)
        if 'shape' in p and (not isinstance(p['shape'], str) or p['shape'] not in {'circle', 'rect', 'pin', 'triangle'}):
            raise ValueError('shape 必须是 circle/rect/pin/triangle')
        if 'orient' in p and (not isinstance(p['orient'], str) or p['orient'] not in {'horizontal', 'vertical'}):
            raise ValueError('orient 必须是 horizontal/vertical')
    return p


class _Output:
    """All basenames are internal constants; no caller-controlled output path."""
    def __init__(self, directory):
        path = Path(directory).expanduser().absolute()
        if any(parent.is_symlink() for parent in [path, *path.parents]):
            raise ValueError('output_dir 不允许 symbolic link')
        path.mkdir(parents=True, exist_ok=True)
        self.path = path.resolve()
        self.paths = []

    def check(self, names):
        for name in names:
            if Path(name).name != name or not re.fullmatch(r'[A-Za-z0-9_-]+\.[a-z0-9]+', name):
                raise ValueError('非法产物名称')
            p = self.path / name
            if p.is_symlink() or (p.exists() and (not p.is_file() or p.stat().st_nlink != 1)):
                raise ValueError('产物不允许 symbolic link 或 hard link')

    def write(self, name, data):
        self.check([name])
        directory_fd = os.open(self.path, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
        temporary_name = '.document-' + uuid.uuid4().hex + '.tmp'
        try:
            fd = os.open(temporary_name, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                0o600, dir_fd=directory_fd)
            with os.fdopen(fd, 'wb') as stream:
                stream.write(data.encode('utf-8') if isinstance(data, str) else data)
            os.replace(temporary_name, name, src_dir_fd=directory_fd, dst_dir_fd=directory_fd)
        finally:
            try:
                os.unlink(temporary_name, dir_fd=directory_fd)
            except FileNotFoundError:
                pass
            os.close(directory_fd)
        path = self.path / name
        self.paths.append(path)
        return path

    def json(self, name, value):
        return self.write(name, json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False))


def _asset_paths(assets):
    if not isinstance(assets, Sequence) or isinstance(assets, (str, bytes)) or len(assets) > 18:
        raise ValueError('assets 必须是最多 18 个已授权文件')
    paths = []
    for asset in assets:
        raw = asset.get('path') if isinstance(asset, Mapping) else asset
        if not isinstance(raw, (str, Path)):
            raise ValueError('素材需要由项目服务解析为 path')
        path = Path(raw)
        if path.is_symlink() or not path.is_file() or path.stat().st_size > 8 * 1024 * 1024:
            raise ValueError('素材文件不可读取、是链接或超过 8 MiB')
        paths.append(path)
    return paths


def _input(text, paths):
    text = _string(text, 'input_text', maximum=100000)
    if not text and paths:
        path = paths[0]
        if path.suffix.lower() not in {'.md', '.txt', '.csv', '.json'}:
            raise ValueError('此操作需要 Markdown/CSV/JSON 文字素材')
        try:
            text = _string(path.read_text('utf-8-sig'), '素材文字', maximum=100000)
        except UnicodeError:
            raise ValueError('素材必须是 UTF-8 编码') from None
    return text


def _load_data(text, p, paths):
    data = p.get('data')
    if data is None:
        raw = _input(text, paths)
        if not raw:
            raise ValueError('需要用户提供真实 CSV/JSON 数据')
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            try:
                reader = csv.DictReader(StringIO(raw))
                if not reader.fieldnames or len(reader.fieldnames) < 2 or len(set(reader.fieldnames)) != len(reader.fieldnames):
                    raise ValueError('CSV 需要不重复的列标题')
                data = list(reader)
                if any(None in row for row in data):
                    raise ValueError('CSV 列数量不一致')
            except csv.Error:
                raise ValueError('数据不是有效 JSON 或 CSV') from None
    try:
        serialized = json.dumps(data, ensure_ascii=False, allow_nan=False)
    except (TypeError, ValueError, OverflowError, RecursionError):
        raise ValueError('data 必须是有效 JSON；所有数字必须有限 / finite') from None
    if len(serialized) > 1_000_000:
        raise ValueError('数据超过 1 MiB')
    return data


def _rows(data, maximum=500):
    if isinstance(data, Mapping) and isinstance(data.get('data'), list):
        data = data['data']
    if not isinstance(data, list) or not data or len(data) > maximum or not all(isinstance(row, Mapping) and row for row in data):
        raise ValueError(f'需要 1–{maximum} 行对象数组数据')
    columns = list(data[0])
    if len(columns) > 30 or not all(isinstance(c, str) and c and len(c) < 100 for c in columns):
        raise ValueError('数据列标题无效或超过 30 列')
    result = []
    for row in data:
        if set(row) != set(columns):
            raise ValueError('数据每行必须使用相同列')
        clean = {}
        for key in columns:
            val = row[key]
            if val is None or val == '':
                clean[key] = None
            elif isinstance(val, (int, float)) and not isinstance(val, bool):
                clean[key] = _number(val, key)
            elif isinstance(val, str):
                value = _string(val, key, maximum=1000)
                try:
                    clean[key] = _number(value, key)
                except ValueError:
                    if value.lower() in {'nan', 'inf', 'infinity', '-inf', '-infinity'}:
                        raise ValueError(f'{key} 必须是有限数字') from None
                    clean[key] = value
            else:
                raise ValueError('表格单元格仅支持文字、有限数字或空值')
        result.append(clean)
    return result


@lru_cache(maxsize=128)
def _font(size, family='sans'):
    from PIL import ImageFont
    candidates = {
        'serif': ['/System/Library/Fonts/Supplemental/Songti.ttc', '/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc'],
        'mono': ['/System/Library/Fonts/Menlo.ttc', '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'],
        'sans': ['/System/Library/Fonts/PingFang.ttc', '/System/Library/Fonts/STHeiti Light.ttc',
            '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',
            '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 'C:/Windows/Fonts/msyh.ttc'],
    }
    for candidate in candidates.get(family, []) + candidates['sans']:
        if Path(candidate).is_file():
            try:
                return ImageFont.truetype(candidate, size=size)
            except OSError:
                continue
    return ImageFont.load_default(size=size)


def _wrap(text, width, size, family='sans'):
    font = _font(size, family)
    # Kinsoku rules apply to soft breaks, while explicit paragraph boundaries
    # remain authored text. Backtrack the break instead of hanging punctuation
    # outside the box or silently deleting a character.
    no_start = frozenset('，。！？：；、）］｝】〕〉》」』〗〙〛”’％‰℃,.:;!?%)]}')
    no_end = frozenset('（［｛【〔〈《「『〖〘〚“‘([{')
    lines = []
    for paragraph in str(text).splitlines() or ['']:
        if not paragraph:
            lines.append('')
            continue
        start = 0
        while start < len(paragraph):
            end = start
            while end < len(paragraph) and font.getlength(paragraph[start:end + 1]) <= width:
                end += 1
            if end == start:
                raise ValueError('文字宽度超过画幅；请增加宽度或减小字号')
            if end < len(paragraph):
                while end > start:
                    before, after = paragraph[start:end].rstrip(), paragraph[end:].lstrip()
                    if not (before and before[-1] in no_end or after and after[0] in no_start):
                        break
                    end -= 1
                if end == start:
                    raise ValueError('画幅太窄，无法保留原文并满足中文换行禁则；请增加宽度或减小字号')
            lines.append(paragraph[start:end])
            start = end
    return lines


def _plain(text):
    return re.sub(r'[*`_]', '', str(text))


class _Scene:
    """One geometric display list feeds both SVG and Pillow rendering."""
    def __init__(self, width, height, bg='#ffffff'):
        self.width, self.height, self.bg = width, height, bg
        self.ops = []

    def rect(self, x, y, w, h, color, outline=None):
        self.ops.append(('rect', (x, y, w, h, color, outline)))

    def line(self, x1, y1, x2, y2, color, thickness=1):
        self.ops.append(('line', (x1, y1, x2, y2, color, thickness)))

    def ellipse(self, x, y, w, h, color):
        self.ops.append(('ellipse', (x, y, w, h, color)))

    def polygon(self, points, color):
        self.ops.append(('polygon', (points, color)))

    def text(self, x, y, text, size, color, family='sans'):
        self.ops.append(('text', (x, y, str(text), int(size), color, family)))

    def paragraph(self, x, y, text, size, color, width, family='sans', line_height=1.45):
        lines = _wrap(text, width, size, family)
        for n, line in enumerate(lines):
            self.text(x, y + n * size * line_height, line, size, color, family)
        return len(lines) * size * line_height

    def svg(self):
        def n(value):
            return f'{value:.3f}'.rstrip('0').rstrip('.') if isinstance(value, float) else str(value)
        tags = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{self.width}" height="{self.height}" viewBox="0 0 {self.width} {self.height}" role="img">',
            f'<rect width="100%" height="100%" fill="{self.bg}"/>']
        for kind, a in self.ops:
            if kind == 'rect':
                x, y, w, h, fill, stroke = a
                tags.append(f'<rect x="{n(x)}" y="{n(y)}" width="{n(w)}" height="{n(h)}" fill="{fill}" stroke="{stroke or "none"}"/>')
            elif kind == 'line':
                x1, y1, x2, y2, color, thickness = a
                tags.append(f'<line x1="{n(x1)}" y1="{n(y1)}" x2="{n(x2)}" y2="{n(y2)}" stroke="{color}" stroke-width="{n(thickness)}"/>')
            elif kind == 'ellipse':
                x, y, w, h, fill = a
                tags.append(f'<ellipse cx="{n(x + w / 2)}" cy="{n(y + h / 2)}" rx="{n(w / 2)}" ry="{n(h / 2)}" fill="{fill}"/>')
            elif kind == 'polygon':
                points, fill = a
                tags.append('<polygon points="' + ' '.join(f'{n(x)},{n(y)}' for x, y in points) + f'" fill="{fill}"/>')
            else:
                x, y, value, size, color, family = a
                stack = {'sans': 'PingFang SC,Noto Sans CJK SC,sans-serif', 'serif': 'Songti SC,Noto Serif CJK SC,serif', 'mono': 'Menlo,Noto Sans Mono CJK SC,monospace'}[family]
                tags.append(f'<text x="{n(x)}" y="{n(y)}" font-size="{size}" font-family="{stack}" dominant-baseline="text-before-edge" fill="{color}">{escape(value)}</text>')
        return ''.join(tags) + '</svg>'

    def image(self, scale=1):
        from PIL import Image, ImageDraw
        img = Image.new('RGB', (self.width * scale, self.height * scale), self.bg)
        draw = ImageDraw.Draw(img)
        for kind, a in self.ops:
            if kind == 'rect':
                x, y, w, h, fill, stroke = a
                draw.rectangle((x * scale, y * scale, (x + w) * scale, (y + h) * scale), fill=fill, outline=stroke, width=scale)
            elif kind == 'line':
                x1, y1, x2, y2, color, thickness = a
                draw.line((x1 * scale, y1 * scale, x2 * scale, y2 * scale), fill=color, width=max(1, round(thickness * scale)))
            elif kind == 'ellipse':
                x, y, w, h, fill = a
                draw.ellipse((x * scale, y * scale, (x + w) * scale, (y + h) * scale), fill=fill)
            elif kind == 'polygon':
                points, fill = a
                draw.polygon([(x * scale, y * scale) for x, y in points], fill=fill)
            else:
                x, y, text, size, color, family = a
                draw.text((x * scale, y * scale), text, font=_font(size * scale, family), fill=color, anchor='lt')
        return img


def _html(title, body, *, interactive=False):
    controls = '<nav><button onclick="zoom(1.2)">+</button><button onclick="zoom(0.8)">−</button><button onclick="zoom(0)">Reset</button></nav>' if interactive else ''
    script = '<script>let z=1;function zoom(f){z=f?Math.min(4,Math.max(.25,z*f)):1;document.querySelector("svg").style.transform="scale("+z+")";}</script>' if interactive else ''
    return '<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">' + \
        f'<title>{escape(title)}</title><style>body{{margin:0;background:#e9e9e7;font-family:system-ui,sans-serif}}svg{{display:block;max-width:100%;height:auto;transform-origin:0 0}}nav{{padding:10px}}button{{padding:8px 16px;margin-right:8px}}main{{overflow:auto}}</style></head><body>{controls}<main>{body}</main>{script}</body></html>'


def _save_scene(out, prefix, scene, p, *, interactive=False, html=True):
    fmt = p.get('format', 'all')
    suffixes = ['svg', 'png'] if fmt == 'all' else [fmt]
    if html and fmt == 'all':
        suffixes.append('html')
    if any(s not in {'svg', 'png', 'html'} for s in suffixes):
        raise ValueError('此渲染操作仅支持 all/html/svg/png')
    out.check([prefix + '.' + ext for ext in suffixes])
    svg = scene.svg()
    if 'svg' in suffixes:
        out.write(prefix + '.svg', svg)
    if 'png' in suffixes:
        buf = BytesIO()
        scene.image(p['scale']).save(buf, format='PNG')
        out.write(prefix + '.png', buf.getvalue())
    if 'html' in suffixes:
        out.write(prefix + '.html', _html(p.get('title', prefix), svg, interactive=interactive))


def _card(text, p, *, number=1, total=1):
    bg, ink, muted, border, accent, family = STYLES[p['style']]
    w, h = p['width'], p['height']
    s = _Scene(w, h, bg)
    landscape = w > h
    pad = max(10, (h if landscape else w) * .056)
    body_size = max(12, round(h / (24 if landscape else 23)))
    header_size = max(24, round(h * .11 if landscape else w * .075))
    lines = [line.strip().lstrip('#').strip() for line in text.splitlines() if line.strip()]
    title = p.get('title') or lines[0]
    body = [line.lstrip('-* ').strip() for line in lines[1:]]
    if p.get('subtitle'):
        body.insert(0, p['subtitle'])
    s.rect(pad, pad, body_size * .5, body_size * .5, accent)
    s.text(pad + body_size, pad - 2, p.get('kicker') or ('QUOTE' if w > h else f'{number:02d} / {total:02d}'), body_size, muted, 'mono')
    s.line(pad, pad + body_size * (1.4 if landscape else 2), w - pad, pad + body_size * (1.4 if landscape else 2), border)
    y = pad + body_size * (1.8 if landscape else 3)
    y += s.paragraph(pad, y, _plain(title), header_size, ink, w - pad * 2, family, 1.3) + body_size
    available = h - pad - body_size * (2.5 if landscape else 3) - y
    wrapped = [(line, _wrap(_plain(line), w - pad * 2 - body_size * 2, body_size, family)) for line in body]
    needed = sum(len(ls) * body_size * 1.45 for _, ls in wrapped)
    if available < 0 or needed > available:
        raise ValueError('卡片文字超出画幅；请拆为多张 cards 或增加画幅')
    gap = max(body_size * .4, (available - needed) / max(1, len(body) - 1))
    for index, (line, _) in enumerate(wrapped):
        s.text(pad, y, f'{index + 1:02d}', body_size, accent, 'mono')
        used = s.paragraph(pad + body_size * 2, y, _plain(line), body_size, ink, w - pad * 2 - body_size * 2, family)
        y += used + gap
        s.line(pad, y - gap / 2, w - pad, y - gap / 2, border)
    footer = p.get('cta') or p.get('source') or p.get('watermark') or (f'{number:02d} / {total:02d}' if not landscape else '')
    s.line(pad, h - pad - body_size * 2, w - pad, h - pad - body_size * 2, border)
    if footer:
        s.paragraph(pad, h - pad - body_size, footer, body_size, muted, w - pad * 2, family)
    return s


def _cards(text, p, out):
    cards = p.get('cards')
    if cards is None:
        if not text:
            raise ValueError('卡片内容为空')
        lines = text.splitlines()
        first_index = next((index for index,line in enumerate(lines) if line.strip()), None)
        if first_index is None:
            raise ValueError('卡片内容为空')
        first = lines[first_index]
        explicit_title = p.get('title')
        # A separately confirmed title replaces a Markdown heading. A plain
        # first paragraph is body content and must not silently disappear.
        body_index = first_index if explicit_title and not re.match(r'^\s{0,3}#{1,6}\s+', first) else first_index + 1
        cards = [{'title': explicit_title or first.lstrip('# ').strip(),
            'body': lines[body_index:]}]
    if not isinstance(cards, list) or not 1 <= len(cards) <= 18:
        raise ValueError('cards 需要 1–18 张卡片')
    scenes = []
    for card in cards:
        if not isinstance(card, Mapping) or set(card) - {'title', 'body'}:
            raise ValueError('每张 card 仅支持 title/body')
        title = _string(card.get('title', ''), 'card.title', maximum=500, empty=False)
        body = card.get('body', [])
        if isinstance(body, str):
            body = body.splitlines()
        if not isinstance(body, list) or len(body) > 30:
            raise ValueError('card.body 必须是最多 30 条文字')
        body = [_string(item, 'card.body', maximum=2000) for item in body]
        scenes.append(_card(title + '\n' + '\n'.join(body), {**p, 'title': title}, number=len(scenes) + 1, total=len(cards)))
    prefixes = ['card'] if len(scenes) == 1 else [f'card_{i}' for i in range(1, len(scenes) + 1)]
    out.check([prefix + '.' + ext for prefix in prefixes for ext in ('svg', 'png', 'html')])
    for prefix, scene in zip(prefixes, scenes):
        _save_scene(out, prefix, scene, p)


def _comparison(data, p, out):
    rows = _rows(data, 20)
    columns = p.get('columns', list(rows[0]))
    if not isinstance(columns, list) or not 2 <= len(columns) <= 4 or not all(c in rows[0] for c in columns):
        raise ValueError('comparison 需要 2–4 个有效 columns')
    bg, ink, muted, border, accent, family = STYLES[p['style']]
    w, h, pad = p['width'], p['height'], p['width'] * .055
    s = _Scene(w, h, bg)
    size = max(12, round(w / 35))
    s.paragraph(pad, pad, p.get('title', ' / '.join(columns[1:])), size * 2, ink, w - pad * 2, family)
    y, cellw = pad + size * 4, (w - pad * 2) / len(columns)
    rowh = max(size * 3, (h - y - pad - size * 3) / (len(rows) + 1))
    if (len(rows) + 1) * rowh > h - y - pad:
        raise ValueError('对比数据超出画幅')
    for ri, row in enumerate([dict(zip(columns, columns)), *rows]):
        for ci, column in enumerate(columns):
            x = pad + ci * cellw
            s.rect(x, y, cellw, rowh, accent if ri == 0 else bg, border)
            lines = _wrap(str(row.get(column, '') or ''), cellw - size, size, family)
            if len(lines) * size * 1.45 > rowh - size:
                raise ValueError('对比单元格文字超出画幅')
            s.paragraph(x + size / 2, y + size / 2, row[column] if row[column] is not None else '', size,
                '#ffffff' if ri == 0 else ink, cellw - size, family)
        y += rowh
    if p.get('source'):
        s.text(pad, h - pad, p['source'], size, muted)
    _save_scene(out, 'comparison', s, p)


def _audit(paths):
    from PIL import Image
    if not paths:
        raise ValueError('audit 需要已授权 PNG/JPEG 素材')
    reports = []
    for path in paths:
        with Image.open(path) as image:
            if image.width * image.height > 16_000_000:
                raise ValueError('audit 素材超过 1600 万像素')
            image = image.convert('RGB')
            w, h = image.size
            if w < 2:
                raise ValueError('audit 图片太窄')
            has = []
            for y in range(h):
                raw = image.crop((0, y, w, y + 1)).tobytes()
                row = list(zip(raw[::3], raw[1::3], raw[2::3]))
                edges = sum(abs(sum(a) - sum(b)) / 3 > 12 for a, b in zip(row, row[1:]))
                has.append(edges / (w - 1) > .006)
            used = [i for i, present in enumerate(has) if present]
            first, last = (used[0], used[-1]) if used else (0, 0)
            gap = current = 0
            for flag in has[first:last + 1]:
                current = 0 if flag else current + 1
                gap = max(gap, current)
            span = (last - first + 1) / h if used else 0
            bottom = (h - last - 1) / h if used else 1
            bands = [sum(has[int(i * h / 4):int((i + 1) * h / 4)]) / max(1, int((i + 1) * h / 4) - int(i * h / 4)) for i in range(4)]
            problems = []
            if bottom > .15:
                problems.append('底部空白超过 15%')
            if gap / h > .15:
                problems.append('中段空白超过 15%')
            if span < .8:
                problems.append('内容跨度不足 80%')
            if bands[3] < .1 and bands[0] > .3:
                problems.append('头重脚轻')
            reports.append({'file': path.name, 'size': f'{w}x{h}', 'span': round(span, 3),
                'bottom_dead': round(bottom, 3), 'max_gap': round(gap / h, 3),
                'fill': round(sum(has) / h, 3), 'bands': [round(v, 2) for v in bands],
                'pass': not problems, 'problems': problems})
    return reports


def _outline(text, title=None):
    raw = [line for line in text.splitlines() if line.strip()]
    if not raw:
        raise ValueError('Markdown 大纲为空')
    first = raw[0].strip()
    first_is_root = bool(re.match(r'^#\s+', first))
    root = {'name': title or first.lstrip('# ').strip(), 'children': []}
    stack = [(0, root)]
    heading_depth = 0
    for line in raw[1:] if first_is_root else raw:
        match = re.match(r'^(#{1,6})\s+(.+)', line.strip())
        if match:
            level, name = len(match[1]) - 1, match[2]
            heading_depth = level
        else:
            indent = len(line) - len(line.lstrip())
            level = heading_depth + 1 + indent // 2
            name = re.sub(r'^\s*(?:[-*+]\s+|\d+[.)]\s+)', '', line).strip()
        while len(stack) > 1 and stack[-1][0] >= level:
            stack.pop()
        node = {'name': _string(name, '大纲节点', maximum=300, empty=False), 'children': []}
        stack[-1][1]['children'].append(node)
        stack.append((level, node))
        if len(stack) > 20:
            raise ValueError('大纲最多 20 层')
    return root


def _tree_scene(root, p):
    if not isinstance(root, Mapping) or set(root) - {'name', 'children', 'description'}:
        raise ValueError('层级数据需要 {name,children?}')
    nodes, edges = [], []
    def visit(node, depth, parent=None):
        if depth > 20 or len(nodes) >= 100:
            raise ValueError('层级最多 100 节点、20 层')
        if set(node) - {'name', 'children', 'description'}:
            raise ValueError('层级节点只支持 name/children/description')
        name = _string(node.get('name'), '节点 name', maximum=300, empty=False)
        if node.get('description'):
            name += '\n' + _string(node['description'], '节点 description', maximum=300)
        index = len(nodes)
        nodes.append((name, depth, 0))
        if parent is not None:
            edges.append((parent, index))
        children = node.get('children', [])
        if not isinstance(children, list) or not all(isinstance(child, Mapping) for child in children):
            raise ValueError('children 必须是节点数组')
        child_indices = [visit(child, depth + 1, index) for child in children]
        if child_indices:
            pos = statistics.mean(nodes[i][2] for i in child_indices)
        else:
            pos = sum(1 for n in nodes[:-1] if n[2] > 0) + 1
        nodes[index] = (name, depth, pos)
        return index
    visit(root, 0)
    levels = max(n[1] for n in nodes) + 1
    leaves = max(n[2] for n in nodes)
    w, h = p['width'], p['height']
    s = _Scene(w, h, _color(p.get('bg', '#ffffff'), 'bg'))
    pad = 30
    vertical = p.get('orient') == 'vertical'
    slotw = (w - pad * 2) / (max(1, leaves) if vertical else levels)
    size = max(11, min(22, round(slotw / 14)))
    positions = [(pad + (pos - .5) / max(1, leaves) * (w - pad * 2) - slotw * .3,
        pad + depth / levels * (h - pad * 2)) if vertical else
        (pad + depth * slotw, pad + (pos - .5) / max(1, leaves) * (h - pad * 2)) for _, depth, pos in nodes]
    for a, b in edges:
        ax, ay = positions[a]
        bx, by = positions[b]
        if vertical:
            s.line(ax + slotw * .3, ay + size * 1.6, bx + slotw * .3, by - 8, _COLORS[b % len(_COLORS)], 2)
        else:
            s.line(ax + slotw * .6, ay + size / 2, bx - 8, by + size / 2, _COLORS[b % len(_COLORS)], 2)
    for i, (name, depth, _) in enumerate(nodes):
        x, y = positions[i]
        lines = _wrap(name, slotw * .8, size)
        if len(lines) * size * 1.45 > (h - pad * 2) / (levels if vertical else max(1, leaves)):
            raise ValueError('层级节点文字超出画幅；请增加尺寸或缩短节点')
        s.paragraph(x, y, name, size, '#24292f', slotw * .8)
        s.line(x, y + len(lines) * size * 1.45 + 3, x + slotw * .65, y + len(lines) * size * 1.45 + 3, _COLORS[i % len(_COLORS)], 2)
    return s


def _chart(rows, p):
    kind = p.get('chart_type', 'column')
    w, h = p['width'], p['height']
    dark = p.get('theme') == 'dark'
    bg, ink, line = ('#0d1117', '#c9d1d9', '#30363d') if dark else ('#ffffff', '#24292f', '#d0d7de')
    s = _Scene(w, h, bg)
    group_key = p.get('group', 'group' if rows and 'group' in rows[0] else None)
    if group_key and kind in {'line', 'area', 'bar', 'column', 'radar'}:
        return _grouped_chart(rows, p, group_key)
    font = max(11, round(w / 40))
    pad = max(36, w * .12)
    top, bottom, right = max(35, font * 3), h - max(40, font * 3), w - max(20, font)
    if p.get('title'):
        s.text(20, 10, p['title'], font * 1.5, ink)
    if kind == 'spreadsheet':
        columns = list(rows[0])
        cellw = (w - 40) / len(columns)
        rowh = (h - top - 20) / (len(rows) + 1)
        if rowh < font * 2:
            raise ValueError('表格数据超出图表画幅')
        for ri, row in enumerate([dict(zip(columns, columns)), *rows]):
            for ci, column in enumerate(columns):
                x, y = 20 + ci * cellw, top + ri * rowh
                s.rect(x, y, cellw, rowh, '#eef4fb' if ri == 0 else bg, line)
                s.paragraph(x + 4, y + 4, row[column] if row[column] is not None else '', font, ink, cellw - 8)
        return s
    if kind == 'scatter':
        xkey, ykey = p.get('x', 'x'), p.get('y', 'y')
        pairs = [(_number(row.get(xkey), xkey), _number(row.get(ykey), ykey)) for row in rows]
        xmin, xmax = min(x for x, _ in pairs), max(x for x, _ in pairs)
        ymin, ymax = min(y for _, y in pairs), max(y for _, y in pairs)
        for x, y in pairs:
            px = pad + (x - xmin) / (xmax - xmin or 1) * (right - pad)
            py = bottom - (y - ymin) / (ymax - ymin or 1) * (bottom - top)
            s.ellipse(px - 4, py - 4, 8, 8, _COLORS[0])
        labels, values = [], []
    else:
        columns = list(rows[0])
        label_key = p.get('x') or next((k for k in ('category', 'time', 'name', 'label') if k in columns), columns[0])
        value_key = p.get('y') or ('value' if 'value' in columns else next((c for c in columns if all(isinstance(row[c], (int, float)) for row in rows)), None))
        if value_key is None:
            raise ValueError('图表需要有限数值列，或用 y 指定数值列')
        labels = [str(row.get(label_key, '')) for row in rows]
        values = [_number(row.get(value_key), value_key) for row in rows]
        if kind in {'pie', 'radar', 'funnel'} and (min(values) < 0 or max(values) <= 0):
            raise ValueError('pie/radar/funnel 需要非负数据且至少一个值大于零')
        if len(rows) > 100:
            raise ValueError('单图最多 100 个数据项')
    if kind == 'histogram':
        bins = p['binNumber']
        lo, hi = min(values), max(values)
        counts = [0] * bins
        for value in values:
            counts[min(bins - 1, int((value - lo) / (hi - lo or 1) * bins))] += 1
        labels = [f'{lo + i * (hi - lo) / bins:.2g}' for i in range(bins)]
        values = counts
        kind = 'column'
    if kind in {'bar', 'column', 'line', 'area', 'histogram', 'scatter'}:
        if kind == 'bar':
            value_width = max(_font(font).getlength(f'{value:g}') for value in values)
            label_width = max(_font(font).getlength(label) for label in labels)
            pad = max(pad, label_width + value_width + 24)
            right = min(right, w - value_width - 18)
            if right - pad < font * 4:
                raise ValueError('条形图标签超出画幅；请缩短分类或增加宽度')
        s.line(pad, top, pad, bottom, line)
        s.line(pad, bottom, right, bottom, line)
        if kind != 'scatter':
            lo, hi = min(0, min(values)), max(0, max(values))
            span = hi - lo or 1
            zero = bottom - (0 - lo) / span * (bottom - top)
            n = len(values)
            if kind == 'bar':
                barh = (bottom - top) / n
                for i, (label, val) in enumerate(zip(labels, values)):
                    left = pad + (min(0, val) - lo) / span * (right - pad)
                    width = abs(val) / span * (right - pad)
                    y = top + i * barh
                    s.rect(left, y + barh * .15, width, barh * .65, _COLORS[i % len(_COLORS)])
                    label_y = y + (barh - font) / 2
                    s.text(4, label_y, label, font, ink)
                    label_text = f'{val:g}'
                    label_x = left - _font(font).getlength(label_text) - 6 if val < 0 else left + width + 6
                    s.text(label_x, label_y, label_text, font, ink)
            elif kind == 'column':
                step = (right - pad) / n
                for i, (label, val) in enumerate(zip(labels, values)):
                    y = bottom - (val - lo) / span * (bottom - top)
                    s.rect(pad + i * step + step * .15, min(y, zero), step * .65, abs(y - zero), _COLORS[i % len(_COLORS)])
                    s.text(pad + i * step, bottom + 8, label, font, ink)
                    s.text(pad + i * step + step * .15, max(top, y - font * 1.4), f'{val:g}', font, ink)
            else:
                points = [(pad + i / max(1, n - 1) * (right - pad), bottom - (val - lo) / span * (bottom - top)) for i, val in enumerate(values)]
                if kind == 'area':
                    s.polygon([(points[0][0], zero), *points, (points[-1][0], zero)], '#cbdcf5')
                for a, b in zip(points, points[1:]):
                    s.line(*a, *b, _COLORS[0], 3)
                for i, (x, y) in enumerate(points):
                    s.ellipse(x - 3, y - 3, 6, 6, _COLORS[0])
                    s.text(x, bottom + 8, labels[i], font, ink)
            for i in range(5):
                value = lo + span * i / 4
                label = f'{value:.3g}'
                if kind == 'bar':
                    xx = pad + i / 4 * (right - pad)
                    s.line(xx, bottom, xx, bottom + 5, line)
                    s.text(xx - _font(font).getlength(label) / 2, bottom + 8, label, font, ink)
                else:
                    yy = bottom - i / 4 * (bottom - top)
                    s.text(4, yy - font / 2, label, font, ink)
        if p.get('axisXTitle'):
            s.text(pad, h - font - 2, p['axisXTitle'], font, ink)
        if p.get('axisYTitle'):
            s.text(2, 4, p['axisYTitle'], font, ink)
    elif kind == 'pie':
        cx, cy = w * .35, (top + bottom) / 2
        radius = min(w * .27, (bottom - top) * .45)
        angle, total = -math.pi / 2, sum(values)
        for i, (label, value) in enumerate(zip(labels, values)):
            end = angle + value / total * math.tau
            samples = max(2, round((end - angle) * 30))
            s.polygon([(cx, cy), *[(cx + radius * math.cos(angle + (end - angle) * j / samples), cy + radius * math.sin(angle + (end - angle) * j / samples)) for j in range(samples + 1)]], _COLORS[i % len(_COLORS)])
            yy = top + i * font * 2
            s.rect(w * .68, yy, font, font, _COLORS[i % len(_COLORS)])
            s.text(w * .68 + font * 1.5, yy, f'{label} {value / total:.1%}', font, ink)
            angle = end
        inner = p.get('innerRadius', 0)
        if inner:
            r = radius * inner
            s.ellipse(cx - r, cy - r, r * 2, r * 2, bg)
    elif kind == 'radar':
        if len(values) < 3:
            raise ValueError('radar 至少需要三个维度')
        cx, cy = w / 2, (top + bottom) / 2
        radius = min(w * .32, (bottom - top) * .4)
        points = []
        for i, (label, value) in enumerate(zip(labels, values)):
            angle = -math.pi / 2 + i / len(values) * math.tau
            ex, ey = cx + radius * math.cos(angle), cy + radius * math.sin(angle)
            s.line(cx, cy, ex, ey, line)
            s.text(ex, ey, label, font, ink)
            points.append((cx + value / max(values) * (ex - cx), cy + value / max(values) * (ey - cy)))
        s.polygon(points, '#cbdcf5')
        for a, b in zip(points, points[1:] + points[:1]):
            s.line(*a, *b, _COLORS[0], 2)
    elif kind == 'funnel':
        step = (bottom - top) / len(values)
        for i, (label, value) in enumerate(zip(labels, values)):
            half = value / max(values) * (w - pad * 2) / 2
            yy = top + i * step
            s.rect(w / 2 - half, yy, half * 2, step * .85, _COLORS[i % len(_COLORS)])
            s.text(pad, yy + 4, f'{label} {value:g}', font, ink)
    return s


def _grouped_chart(rows, p, group_key):
    """Keep series independent and share one numeric axis across them."""
    kind, w, h = p['chart_type'], p['width'], p['height']
    columns = list(rows[0])
    xkey = p.get('x') or next((key for key in ('time', 'category', 'name', 'label') if key in columns), columns[0])
    ykey = p.get('y', 'value')
    groups, labels = {}, []
    for row in rows:
        if group_key not in row or row[group_key] is None or row.get(xkey) is None:
            raise ValueError('group 和分类列不能为空')
        group, label = str(row[group_key]), str(row[xkey])
        if label not in labels:
            labels.append(label)
        points = groups.setdefault(group, {})
        if label in points:
            raise ValueError('同一 group/分类重复，需先聚合数据')
        points[label] = _number(row.get(ykey), ykey)
    if len(groups) > 20 or len(labels) > 100:
        raise ValueError('图表最多 20 系列、100 分类')
    values = [value for series in groups.values() for value in series.values()]
    lo, hi = min(0, min(values)), max(0, max(values))
    s = _Scene(w, h, '#0d1117' if p.get('theme') == 'dark' else '#ffffff')
    ink = '#c9d1d9' if p.get('theme') == 'dark' else '#24292f'
    left, right, top, bottom = w * .12, w - 30, 50, h - 45
    font = max(11, round(w / 45))
    s.line(left, top, left, bottom, '#d0d7de'); s.line(left, bottom, right, bottom, '#d0d7de')
    if p.get('title'):
        s.text(15, 8, p['title'], font * 1.5, ink)
    for gi, (group, vals) in enumerate(groups.items()):
        color = _COLORS[gi % len(_COLORS)]
        s.text(15 + gi * (w - 30) / len(groups), h - font - 2, group, font, color)
        if kind == 'radar':
            if min(values) < 0 or hi <= 0 or len(labels) < 3:
                raise ValueError('分组 radar 需要至少三个维度和非负数值')
            cx, cy, radius = w / 2, (top + bottom) / 2, min(w, h) * .32
            points = []
            for i, label in enumerate(labels):
                if label not in vals:
                    raise ValueError('radar 每组必须包含相同维度')
                angle = -math.pi / 2 + i / len(labels) * math.tau
                ex, ey = cx + radius * math.cos(angle), cy + radius * math.sin(angle)
                if gi == 0:
                    s.line(cx, cy, ex, ey, '#d0d7de'); s.text(ex, ey, label, font, ink)
                points.append((cx + (ex - cx) * vals[label] / hi, cy + (ey - cy) * vals[label] / hi))
            for a, b in zip(points, points[1:] + points[:1]):
                s.line(*a, *b, color, 2)
        elif kind in {'bar', 'column'}:
            step = ((bottom - top) if kind == 'bar' else (right - left)) / len(labels)
            for i, label in enumerate(labels):
                if label not in vals:
                    continue
                value = vals[label]
                if kind == 'column':
                    x = left + i * step + step * .08 + gi * step * .8 / len(groups)
                    y = bottom - (value - lo) / (hi - lo or 1) * (bottom - top)
                    zero = bottom - (0 - lo) / (hi - lo or 1) * (bottom - top)
                    s.rect(x, min(y, zero), step * .7 / len(groups), abs(y - zero), color)
                    if gi == 0:
                        s.text(left + i * step, bottom + 5, label, font, ink)
                else:
                    x = left + (min(0, value) - lo) / (hi - lo or 1) * (right - left)
                    y = top + i * step + step * .08 + gi * step * .8 / len(groups)
                    s.rect(x, y, abs(value) / (hi - lo or 1) * (right - left), step * .7 / len(groups), color)
                    if gi == 0:
                        s.text(3, top + i * step, label, font, ink)
        else:
            run = []
            def draw_run(points):
                if len(points) > 1:
                    if kind == 'area':
                        zero = bottom - (0 - lo) / (hi - lo or 1) * (bottom - top)
                        # Pale fill is distinct from the colored series outline.
                        s.polygon([(points[0][0], zero), *points, (points[-1][0], zero)], '#e5edf9' if gi == 0 else '#e8f7f1')
                    for a, b in zip(points, points[1:]):
                        s.line(*a, *b, color, 3)
                for x, y in points:
                    s.ellipse(x - 3, y - 3, 6, 6, color)
            for i, label in enumerate(labels):
                if label not in vals:
                    draw_run(run); run = []
                    continue
                run.append((left + i / max(1, len(labels) - 1) * (right - left), bottom - (vals[label] - lo) / (hi - lo or 1) * (bottom - top)))
                if gi == 0:
                    s.text(left + i / max(1, len(labels) - 1) * (right - left), bottom + 5, label, font, ink)
            draw_run(run)
    for i in range(5):
        s.text(3, bottom - i / 4 * (bottom - top), f'{lo + (hi - lo) * i / 4:.3g}', font, ink)
    return s


def _statistical_chart(rows, p):
    kind, w, h = p['chart_type'], p['width'], p['height']
    s = _Scene(w, h)
    groups = {}
    for row in rows:
        label = str(row.get(p.get('x', 'category'), ''))
        groups.setdefault(label, []).append(_number(row.get(p.get('y', 'value')), 'value'))
    if len(groups) > 30:
        raise ValueError('统计图最多 30 个类别')
    low, high = min(min(v) for v in groups.values()), max(max(v) for v in groups.values())
    span = high - low or 1
    left, right, top, bottom = 50, w - 25, 45, h - 40
    font = max(11, round(w / 45))
    s.line(left, top, left, bottom, '#d0d7de'); s.line(left, bottom, right, bottom, '#d0d7de')
    if p.get('title'):
        s.text(15, 10, p['title'], font * 1.5, '#24292f')
    def yy(value):
        return bottom - (value - low) / span * (bottom - top)
    slot = (right - left) / len(groups)
    for i, (label, vals) in enumerate(groups.items()):
        cx, half = left + (i + .5) * slot, slot * .28
        color = _COLORS[i % len(_COLORS)]
        ordered = sorted(vals)
        q1, median, q3 = statistics.quantiles(ordered, n=4, method='inclusive') if len(vals) > 1 else (vals[0], vals[0], vals[0])
        if kind == 'boxplot':
            fence_low, fence_high = q1 - 1.5 * (q3 - q1), q3 + 1.5 * (q3 - q1)
            inliers = [value for value in vals if fence_low <= value <= fence_high]
            bottom_whisker, top_whisker = min(inliers), max(inliers)
            s.line(cx, yy(bottom_whisker), cx, yy(top_whisker), color, 2)
            s.line(cx - half / 2, yy(bottom_whisker), cx + half / 2, yy(bottom_whisker), color, 2)
            s.line(cx - half / 2, yy(top_whisker), cx + half / 2, yy(top_whisker), color, 2)
            s.rect(cx - half, yy(q3), half * 2, max(1, yy(q1) - yy(q3)), '#e5edf9', color)
            s.line(cx - half, yy(median), cx + half, yy(median), color, 3)
            for val in vals:
                if val < fence_low or val > fence_high:
                    s.ellipse(cx - 3, yy(val) - 3, 6, 6, color)
        else:
            # Gaussian KDE with Scott's bandwidth; computation uses real samples.
            std = statistics.stdev(vals) if len(vals) > 1 else 0
            bandwidth = max(span / 100, std * len(vals) ** -.2, 1e-8)
            grid = [low + span * j / 60 for j in range(61)]
            densities = [sum(math.exp(-.5 * ((point - val) / bandwidth) ** 2) for val in vals) / (len(vals) * bandwidth * math.sqrt(math.tau)) for point in grid]
            peak = max(densities) or 1
            right_points = [(cx + half * density / peak, yy(value)) for value, density in zip(grid, densities)]
            left_points = [(cx - half * density / peak, yy(value)) for value, density in reversed(list(zip(grid, densities)))]
            s.polygon(right_points + left_points, color)
            s.line(cx - half / 2, yy(median), cx + half / 2, yy(median), '#ffffff', 2)
        s.text(cx - half, bottom + 8, label, font, '#24292f')
    return s


def _waterfall(data, p):
    if not isinstance(data, list) or not 1 <= len(data) <= 100:
        raise ValueError('waterfall 需要 1–100 个数据项')
    stages, running = [], 0
    for row in data:
        if not isinstance(row, Mapping) or set(row) - {'category', 'value', 'isTotal', 'isIntermediateTotal'}:
            raise ValueError('waterfall 数据字段无效')
        label = _string(row.get('category'), 'category', maximum=100, empty=False)
        for key in ('isTotal', 'isIntermediateTotal'):
            if key in row and not isinstance(row[key], bool):
                raise ValueError(f'{key} 必须是布尔值')
        if row.get('isTotal') or row.get('isIntermediateTotal'):
            finish = _number(row['value'], 'value') if 'value' in row else running
            stages.append((label, 0, finish, 'total'))
            running = finish
        else:
            value = _number(row.get('value'), 'value')
            stages.append((label, running, running + value, 'positive' if value >= 0 else 'negative'))
            running += value
    w, h = p['width'], p['height']
    s = _Scene(w, h)
    lo, hi = min(0, min(min(start, end) for _, start, end, _ in stages)), max(0, max(max(start, end) for _, start, end, _ in stages))
    left, right, top, bottom = 45, w - 25, 40, h - 45
    step, font = (right - left) / len(stages), max(11, round(w / 45))
    def yy(value):
        return bottom - (value - lo) / (hi - lo or 1) * (bottom - top)
    for i, (label, start, end, state) in enumerate(stages):
        x, y1, y2 = left + i * step, yy(start), yy(end)
        color = {'total': '#5B8FF9', 'positive': '#61DDAA', 'negative': '#E86452'}[state]
        s.rect(x + step * .1, min(y1, y2), step * .75, max(1, abs(y1 - y2)), color)
        s.text(x, bottom + 5, label, font, '#24292f')
        s.text(x, max(5, min(y1, y2) - font * 1.4), f'{end if state == "total" else end - start:g}', font, '#24292f')
        if i < len(stages) - 1:
            s.line(x + step * .85, y2, x + step * 1.1, y2, '#d0d7de')
    s.line(left, yy(0), right, yy(0), '#d0d7de')
    return s


def _dual_axes(data, p):
    if not isinstance(data, Mapping) or not isinstance(data.get('categories'), list) or not isinstance(data.get('series'), list):
        raise ValueError('dual-axes 需要 {categories,series}')
    labels, series = data['categories'], data['series']
    if not 1 <= len(labels) <= 100 or not 1 <= len(series) <= 2:
        raise ValueError('dual-axes 需要 1–100 分类、1–2 系列')
    labels = [_string(str(label), 'category', maximum=100) for label in labels]
    w, h = p['width'], p['height']
    s = _Scene(w, h)
    left, right, top, bottom = 50, w - 50, 45, h - 45
    font = max(11, round(w / 45))
    s.line(left, top, left, bottom, '#d0d7de'); s.line(right, top, right, bottom, '#d0d7de')
    for si, item in enumerate(series):
        if not isinstance(item, Mapping) or item.get('type') not in {'line', 'column'} or not isinstance(item.get('data'), list) or len(item['data']) != len(labels):
            raise ValueError('每条 series 需要 type=line/column 及等长 data')
        vals = [_number(value, 'series value') for value in item['data']]
        lo, hi = min(0, min(vals)), max(0, max(vals))
        color = _COLORS[si]
        points = [(left + (i + .5) / len(labels) * (right - left), bottom - (val - lo) / (hi - lo or 1) * (bottom - top)) for i, val in enumerate(vals)]
        if item['type'] == 'line':
            for a, b in zip(points, points[1:]):
                s.line(*a, *b, color, 3)
            for x, y in points:
                s.ellipse(x - 3, y - 3, 6, 6, color)
        else:
            zero = bottom - (0 - lo) / (hi - lo or 1) * (bottom - top)
            half = (right - left) / len(labels) * .27
            for x, y in points:
                s.rect(x - half, min(y, zero), half * 2, abs(y - zero), color)
        for i in range(5):
            s.text(3 if si == 0 else right + 5, bottom - i / 4 * (bottom - top), f'{lo + (hi - lo) * i / 4:.3g}', font, color)
        if item.get('axisYTitle'):
            title = _string(item['axisYTitle'], 'axisYTitle', maximum=100)
            s.text(left if si == 0 else right - len(title) * font / 2, 12, title, font, color)
    for i, label in enumerate(labels):
        s.text(left + (i + .2) / len(labels) * (right - left), bottom + 8, label, font, '#24292f')
    return s


def _graph(data, p):
    if not isinstance(data, Mapping) or not isinstance(data.get('nodes'), list) or not isinstance(data.get('edges'), list):
        raise ValueError('关系/流程图需要 {nodes,edges}')
    nodes, edges = data['nodes'], data['edges']
    if not 1 <= len(nodes) <= 60 or len(edges) > 200 or not all(isinstance(node, Mapping) for node in nodes):
        raise ValueError('关系图最多 60 节点、200 条边')
    names = [_string(node.get('name'), 'node.name', maximum=100, empty=False) for node in nodes]
    if len(set(names)) != len(names):
        raise ValueError('节点名称不能重复')
    adjacency, indegree, normalized = {name: [] for name in names}, dict.fromkeys(names, 0), []
    for edge in edges:
        if not isinstance(edge, Mapping) or edge.get('source') not in names or edge.get('target') not in names:
            raise ValueError('edge 引用了不存在的节点')
        source, target = edge['source'], edge['target']
        label = _string(edge.get('name', ''), 'edge.name', maximum=100)
        adjacency[source].append(target); indegree[target] += 1
        normalized.append((source, target, label))
    w, h = p['width'], p['height']
    s = _Scene(w, h)
    font = max(11, round(w / 45))
    if p['chart_type'] == 'flow-diagram':
        queue = [name for name in names if not indegree[name]]
        depths = dict.fromkeys(queue, 0)
        for name in queue:
            for target in adjacency[name]:
                depths[target] = max(depths.get(target, 0), depths[name] + 1)
                indegree[target] -= 1
                if indegree[target] == 0:
                    queue.append(target)
        if len(queue) != len(names):
            raise ValueError('flow-diagram 需要无环流程；循环关系请用 network-graph')
        levels = max(depths.values()) + 1
        by_depth = {depth: [name for name in names if depths[name] == depth] for depth in range(levels)}
        positions = {name: (40 + (depth + .5) / levels * (w - 80), 45 + (i + .5) / len(layer) * (h - 90)) for depth, layer in by_depth.items() for i, name in enumerate(layer)}
    else:
        radius = min(w, h) * .34
        positions = {name: (w / 2 + radius * math.cos(-math.pi / 2 + i / len(names) * math.tau), h / 2 + radius * math.sin(-math.pi / 2 + i / len(names) * math.tau)) for i, name in enumerate(names)}
    for source, target, label in normalized:
        ax, ay = positions[source]; bx, by = positions[target]
        s.line(ax, ay, bx, by, '#737373', 2)
        dx, dy = bx - ax, by - ay
        length = math.hypot(dx, dy)
        if length:
            ux, uy = dx / length, dy / length
            tipx, tipy = bx - ux * font * 2, by - uy * font * 1.5
            s.polygon([(tipx, tipy), (tipx - ux * 9 - uy * 4, tipy - uy * 9 + ux * 4), (tipx - ux * 9 + uy * 4, tipy - uy * 9 - ux * 4)], '#737373')
        if label:
            s.text((ax + bx) / 2, (ay + by) / 2, label, font, '#24292f')
    for i, name in enumerate(names):
        cx, cy = positions[name]
        cellw = max(font * 4, _font(font).getlength(name) + 12)
        s.rect(cx - cellw / 2, cy - font, cellw, font * 2.4, '#f6f8fa', _COLORS[i % len(_COLORS)])
        s.text(cx - cellw / 2 + 6, cy - font / 2, name, font, '#24292f')
    return s


def _sankey(data, p):
    if not isinstance(data, list) or not 1 <= len(data) <= 100:
        raise ValueError('sankey 需要 1–100 条流量')
    flows, names = [], []
    for row in data:
        if not isinstance(row, Mapping):
            raise ValueError('sankey 每条流量必须是对象')
        source = _string(row.get('source'), 'source', maximum=100, empty=False)
        target = _string(row.get('target'), 'target', maximum=100, empty=False)
        value = _number(row.get('value'), 'value', .00000001)
        if source == target:
            raise ValueError('sankey 不支持自环')
        flows.append((source, target, value))
        for name in (source, target):
            if name not in names:
                names.append(name)
    graph_data = {'nodes': [{'name': name} for name in names], 'edges': [{'source': source, 'target': target} for source, target, _ in flows]}
    # Reuse DAG admission and layout, then replace thin edges by weighted ribbons.
    _graph(graph_data, {**p, 'chart_type': 'flow-diagram'})
    indegree = {name: sum(target == name for _, target, _ in flows) for name in names}
    queue = [name for name in names if indegree[name] == 0]
    depths = dict.fromkeys(queue, 0)
    for name in queue:
        for source, target, _ in flows:
            if source == name:
                depths[target] = max(depths.get(target, 0), depths[name] + 1)
                indegree[target] -= 1
                if indegree[target] == 0:
                    queue.append(target)
    w, h = p['width'], p['height']
    s = _Scene(w, h)
    font = max(11, round(w / 45))
    levels = max(depths.values()) + 1
    weights = {name: max(sum(value for source, _, value in flows if source == name), sum(value for _, target, value in flows if target == name)) for name in names}
    by_depth = {depth: [name for name in names if depths[name] == depth] for depth in range(levels)}
    unit = min((h - 90 - 15 * max(0, len(layer) - 1)) / sum(weights[name] for name in layer) for layer in by_depth.values())
    if unit <= 0:
        raise ValueError('sankey 节点过多，画幅不足')
    positions = {}
    for depth, layer in by_depth.items():
        y = 45
        for name in layer:
            positions[name] = (30 + depth / max(1, levels - 1) * (w - 70), y, weights[name] * unit)
            y += weights[name] * unit + 15
    offsets_in, offsets_out = dict.fromkeys(names, 0), dict.fromkeys(names, 0)
    for i, (source, target, value) in enumerate(flows):
        sx, sy, _ = positions[source]; tx, ty, _ = positions[target]
        sy += offsets_out[source]; ty += offsets_in[target]
        ribbon = value * unit
        s.polygon([(sx + 14, sy), (tx, ty), (tx, ty + ribbon), (sx + 14, sy + ribbon)], _COLORS[i % len(_COLORS)])
        offsets_out[source] += ribbon; offsets_in[target] += ribbon
    for name, (x, y, nodeh) in positions.items():
        s.rect(x, y, 14, nodeh, '#24292f')
        s.text(max(0, x - font), y - font - 3, f'{name} {weights[name]:g}', font, '#24292f')
    return s


def _treemap(data, p):
    if not isinstance(data, list) or not data or len(data) > 100:
        raise ValueError('treemap 需要非空节点数组')
    total_nodes = 0
    def clean(nodes, depth=0):
        nonlocal total_nodes
        if depth > 3:
            raise ValueError('treemap 最深三层')
        result = []
        for node in nodes:
            if not isinstance(node, Mapping):
                raise ValueError('treemap 节点必须是对象')
            name = _string(node.get('name'), 'name', maximum=100, empty=False)
            children = node.get('children', [])
            if not isinstance(children, list):
                raise ValueError('treemap children 必须是数组')
            children = clean(children, depth + 1) if children else []
            value = _number(node.get('value', sum(c['value'] for c in children)), 'value', .00000001)
            total_nodes += 1
            if total_nodes > 100:
                raise ValueError('treemap 最多 100 节点')
            result.append({'name': name, 'value': value, 'children': children})
        return result
    nodes = clean(data)
    s = _Scene(p['width'], p['height'])
    font = max(10, round(p['width'] / 50))
    def layout(items, x, y, width, height, depth=0):
        total = sum(node['value'] for node in items)
        for i, node in enumerate(items):
            fraction = node['value'] / total
            cw, ch = (width * fraction, height) if depth % 2 == 0 else (width, height * fraction)
            if min(cw, ch) < font * 2:
                raise ValueError('treemap 有占比过小的节点，无法在当前画幅可读显示')
            color = _COLORS[(i + depth * 2) % len(_COLORS)]
            s.rect(x + 1, y + 1, cw - 2, ch - 2, color, '#ffffff')
            s.paragraph(x + 5, y + 5, f'{node["name"]} {node["value"]:g}', font, '#24292f', cw - 10)
            if node['children']:
                layout(node['children'], x + 5, y + font * 2, cw - 10, ch - font * 2 - 5, depth + 1)
            if depth % 2 == 0:
                x += cw
            else:
                y += ch
    layout(nodes, 20, 35, p['width'] - 40, p['height'] - 55)
    return s


def _wordcloud(rows, p):
    if len(rows) > 100:
        raise ValueError('词云最多 100 个词')
    items = [(_string(row.get('text'), 'text', maximum=100, empty=False), _number(row.get('value'), 'value', 0)) for row in rows]
    maximum = max(value for _, value in items)
    if maximum <= 0:
        raise ValueError('词云至少一个权重大于零')
    w, h = p['width'], p['height']
    s, occupied = _Scene(w, h), []
    for index, (word, value) in enumerate(sorted(items, key=lambda item: item[1], reverse=True)):
        size = round(12 + math.sqrt(value / maximum) * min(w, h) * .13)
        placed = None
        while size >= 10 and placed is None:
            fw, fh = _font(size).getlength(word) + 6, size * 1.4
            for step in range(2500):
                angle, radius = step * .28, math.sqrt(step) * min(w, h) / 110
                x, y = w / 2 + math.cos(angle) * radius - fw / 2, h / 2 + math.sin(angle) * radius - fh / 2
                candidate = (x, y, x + fw, y + fh)
                if x < 5 or y < 5 or x + fw > w - 5 or y + fh > h - 5:
                    continue
                if not any(candidate[0] < rect[2] and candidate[2] > rect[0] and candidate[1] < rect[3] and candidate[3] > rect[1] for rect in occupied):
                    placed = candidate; break
            if placed is None:
                size -= 2
        if placed is None:
            raise ValueError('词云文字无法完整排入画幅，请增加尺寸或减少词数量')
        occupied.append(placed)
        s.text(placed[0] + 3, placed[1] + 3, word, size, _COLORS[index % len(_COLORS)])
    return s


def _venn(data, p):
    if not isinstance(data, list) or not data or len(data) > 7:
        raise ValueError('venn 需要最多三个集合的七个区域')
    regions, names = {}, []
    for row in data:
        if not isinstance(row, Mapping) or not isinstance(row.get('sets'), list) or not 1 <= len(row['sets']) <= 3:
            raise ValueError('venn 需要 {sets:[名称],value}')
        sets = [_string(name, 'set', maximum=100, empty=False) for name in row['sets']]
        if len(set(sets)) != len(sets):
            raise ValueError('sets 不能重复')
        key = frozenset(sets)
        if key in regions:
            raise ValueError('venn 区域重复')
        regions[key] = _number(row.get('value'), 'value', 0)
        for name in sets:
            if name not in names:
                names.append(name)
    if not 1 <= len(names) <= 3:
        raise ValueError('venn 最多三个集合')
    for sets, value in regions.items():
        if len(sets) > 1 and any(frozenset([name]) in regions and value > regions[frozenset([name])] for name in sets):
            raise ValueError('集合交集不能大于单个集合')
    w, h = p['width'], p['height']
    s = _Scene(w, h)
    radius = min(w / (3.4 if len(names) > 1 else 2.4), h / 2.7)
    centers = [(w / 2, h / 2)] if len(names) == 1 else [(w / 2 - radius * .5, h / 2), (w / 2 + radius * .5, h / 2)]
    if len(names) == 3:
        centers = [(w / 2 - radius * .5, h * .42), (w / 2 + radius * .5, h * .42), (w / 2, h * .62)]
    font = max(11, round(w / 45))
    # Membership schematic: region counts are exact; circle areas aren't scaled.
    for i, (name, (cx, cy)) in enumerate(zip(names, centers)):
        color = ['#d9e7fb', '#d8f2e5', '#faeccb'][i]
        s.ellipse(cx - radius, cy - radius, radius * 2, radius * 2, color)
        s.text(cx - radius * .6, cy - radius * .75, name, font * 1.2, '#24292f')
    for key, value in regions.items():
        indices = [names.index(name) for name in key]
        cx = statistics.mean(centers[i][0] for i in indices)
        cy = statistics.mean(centers[i][1] for i in indices)
        if len(key) == 1 and len(names) > 1:
            cx += (-1 if indices[0] == 0 else 1 if indices[0] == 1 else 0) * radius * .5
            cy += radius * .35 if indices[0] == 2 else -radius * .05
        s.text(cx - font, cy, f'{value:g}', font * 1.4, '#24292f')
    s.text(10, h - font * 1.5, '集合关系示意；圆面积不代表数量', font, '#737373')
    return s


def _fishbone(data, p):
    # The same tree admission prevents malformed/deep input before rendering.
    _tree_scene(data, p)
    w, h = p['width'], p['height']
    s = _Scene(w, h)
    font, mid = max(11, round(w / 45)), h / 2
    children = data.get('children', [])
    s.line(20, mid, w - 130, mid, '#737373', 3)
    s.polygon([(w - 140, mid - 7), (w - 126, mid), (w - 140, mid + 7)], '#737373')
    s.paragraph(w - 120, mid - font, data['name'], font * 1.4, '#24292f', 110)
    for i, child in enumerate(children):
        x = 40 + (i + .5) / max(1, len(children)) * (w - 200)
        y = mid + (-1 if i % 2 == 0 else 1) * h * .3
        s.line(x - 30, y, x + 30, mid, _COLORS[i % len(_COLORS)], 2)
        s.text(x - 50, y - font, child['name'], font, '#24292f')
        for j, sub in enumerate(child.get('children', [])):
            sy = y + (mid - y) * (j + 1) / (len(child['children']) + 1)
            s.line(x - 45, sy, x, sy, '#737373')
            s.text(x - 65, sy - font, sub['name'], font, '#24292f')
    return s


def _prepare_chart(data, p):
    kind = p['chart_type']
    if kind in {'mind-map', 'organization-chart'}:
        return _tree_scene(data, p), data
    if kind == 'fishbone-diagram':
        return _fishbone(data, p), data
    if kind in {'network-graph', 'flow-diagram'}:
        return _graph(data, p), data
    if kind == 'sankey':
        return _sankey(data, p), data
    if kind == 'treemap':
        return _treemap(data, p), data
    if kind == 'venn':
        return _venn(data, p), data
    if kind == 'waterfall':
        return _waterfall(data, p), data
    if kind == 'dual-axes':
        return _dual_axes(data, p), data
    if kind == 'liquid':
        if not isinstance(data, Mapping):
            raise ValueError('liquid 需要 {percent:0–1}')
        percent = _number(data.get('percent'), 'percent', 0, 1)
        w, h = p['width'], p['height']
        s = _Scene(w, h)
        cx, cy, radius = w / 2, h / 2, min(w, h) * .38
        shape = p.get('shape', 'circle')
        if shape in {'circle', 'pin'}:
            s.ellipse(cx - radius, cy - radius, radius * 2, radius * 2, '#e9edf2')
            water = cy + radius - 2 * radius * percent
            arc = [(cx + radius * math.cos(t), cy + radius * math.sin(t)) for t in [i / 160 * math.tau for i in range(161)] if cy + radius * math.sin(t) >= water]
            if arc:
                arc.sort(key=lambda point: math.atan2(point[1] - cy, point[0] - cx))
                s.polygon(arc, _COLORS[0])
            if shape == 'pin':
                s.polygon([(cx - radius * .4, cy + radius * .75), (cx, cy + radius * 1.25), (cx + radius * .4, cy + radius * .75)], _COLORS[0])
        elif shape == 'triangle':
            s.polygon([(cx, cy - radius), (cx + radius, cy + radius), (cx - radius, cy + radius)], '#e9edf2')
            half = radius * (1 - percent)
            water = cy + radius - 2 * radius * percent
            s.polygon([(cx - radius, cy + radius), (cx + radius, cy + radius), (cx + half, water), (cx - half, water)], _COLORS[0])
        else:
            s.rect(cx - radius, cy - radius, radius * 2, radius * 2, '#e9edf2')
            s.rect(cx - radius, cy + radius - radius * 2 * percent, radius * 2, radius * 2 * percent, _COLORS[0])
        font = max(20, round(min(w, h) * .1))
        s.text(cx - font, cy - font / 2, f'{percent:.0%}', font, '#24292f')
        return s, data
    if kind == 'histogram' and isinstance(data, list) and data and all(isinstance(v, (int, float)) for v in data):
        data = [{'category': str(i), 'value': _number(v, 'value')} for i, v in enumerate(data)]
    rows = _rows(data)
    if kind in {'boxplot', 'violin'}:
        return _statistical_chart(rows, p), rows
    if kind == 'word-cloud':
        return _wordcloud(rows, p), rows
    return _chart(rows, p), rows


def _blocks(text):
    lines, blocks, i = text.splitlines(), [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith('```'):
            language, code = line.strip()[3:], []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i += 1
            blocks.append(('code', '\n'.join(code), language)); i += 1
            continue
        if i + 1 < len(lines) and '|' in line and re.fullmatch(r'[\s|:-]+', lines[i + 1]) and '-' in lines[i + 1]:
            rows = [[cell.strip() for cell in line.strip().strip('|').split('|')]]
            i += 2
            while i < len(lines) and '|' in lines[i] and lines[i].strip():
                rows.append([cell.strip() for cell in lines[i].strip().strip('|').split('|')]); i += 1
            if any(len(row) != len(rows[0]) for row in rows):
                raise ValueError('Markdown 表格列数不一致')
            blocks.append(('table', rows, None)); continue
        heading = re.match(r'^(#{1,6})\s+(.+)$', line.strip())
        if heading:
            blocks.append(('heading', heading[2], len(heading[1])))
        elif re.match(r'^\s*(?:[-*+]\s+|\d+[.)]\s+)', line):
            blocks.append(('list', re.sub(r'^\s*(?:[-*+]\s+|\d+[.)]\s+)', '', line), None))
        elif line.lstrip().startswith('>'):
            blocks.append(('quote', line.lstrip()[1:].strip(), None))
        elif re.fullmatch(r'\s*(?:---+|\*\*\*+)\s*', line):
            blocks.append(('rule', '', None))
        else:
            paragraph = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^\s*(?:#|[-*+]\s|\d+[.)]\s|>|```|\|)', lines[i + 1]):
                i += 1; paragraph.append(lines[i])
            blocks.append(('paragraph', '\n'.join(paragraph), None))
        i += 1
    return blocks


def _inline(text):
    value = escape(str(text))
    value = re.sub(r'`([^`]+)`', r'<code>\1</code>', value)
    value = re.sub(r'\*\*([^*]+)\*\*', r'<strong>\1</strong>', value)
    value = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'<em>\1</em>', value)
    return value.replace('\n', '<br>')


def _document_html(blocks, title, p, *, gzh=False):
    theme = p.get('theme', 'moyu-green')
    accent = GZH_THEMES[theme] if gzh else '#0969da'
    body = []
    for kind, value, meta in blocks:
        if kind == 'heading':
            size = [30, 25, 21, 19, 18, 17][meta - 1]
            body.append(f'<h{meta} style="font-size:{size}px;line-height:1.35;color:{accent};border-bottom:1px solid #d0d7de;padding-bottom:8px;margin:24px 0 12px">{_inline(value)}</h{meta}>')
        elif kind == 'code':
            code = f'<pre style="background:#f6f8fa;padding:16px;border-radius:6px;white-space:pre-wrap"><code>{escape(value)}</code></pre>'
            if gzh:
                code = code.replace('<code>', '<span style="font-family:monospace">').replace('</code>', '</span>')
            body.append(code)
        elif kind == 'table':
            if gzh:
                body.append('<section style="border:1px solid #d0d7de">' + ''.join(
                    '<section style="display:flex">' + ''.join(f'<span style="flex:1;padding:8px;border-bottom:1px solid #d0d7de">{_inline(cell)}</span>' for cell in row) + '</section>' for row in value) + '</section>')
            else:
                body.append('<table><thead><tr>' + ''.join(f'<th>{_inline(cell)}</th>' for cell in value[0]) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join(f'<td>{_inline(cell)}</td>' for cell in row) + '</tr>' for row in value[1:]) + '</tbody></table>')
        elif kind == 'rule':
            body.append('<hr style="border:0;border-top:1px solid #d0d7de">')
        else:
            prefix = '• ' if kind == 'list' else ''
            border = f'border-left:4px solid {accent};padding-left:16px;' if kind == 'quote' else ''
            content = _inline(prefix + value)
            if gzh:
                content = content.replace('<code>', '<span style="font-family:monospace;background:#eeefe9;padding:2px 4px">').replace('</code>', '</span>')
            body.append(f'<p style="margin:12px 0;line-height:1.75;{border}">{content}</p>')
    width = p.get('page_width', 677 if gzh else 760)
    section = f'<section style="max-width:{width}px;margin:0 auto;padding:32px;color:#24292f;background:#ffffff;font:17px/1.75 -apple-system,PingFang SC,Noto Sans CJK SC,sans-serif">' + ''.join(body) + '</section>'
    if gzh:
        section = f'<!-- theme: {theme} -->' + section + '<p style="display:none"><mp-style-type data-value="3"></mp-style-type></p>'
        return '<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"><title>' + escape(title) + '</title></head><body>' + section + '</body></html>'
    css = 'body{margin:0;background:#fff}table{border-collapse:collapse;width:100%}th,td{border:1px solid #d0d7de;padding:8px 12px;text-align:left}th{background:#f6f8fa}code{background:#f6f8fa;padding:.15em .4em;border-radius:5px}pre code{padding:0}'
    return '<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + escape(title) + '</title><style>' + css + '</style></head><body>' + section + '</body></html>'


def _document_scene(blocks, p):
    w = p['width']
    pad, size, y = max(20, w * .05), max(14, round(w / 47)), max(20, w * .05)
    s = _Scene(w, 1)
    for kind, value, meta in blocks:
        fs = round(size * (1.9 if meta == 1 else 1.4)) if kind == 'heading' else size
        if kind == 'table':
            cellw = (w - pad * 2) / len(value[0])
            for ri, row in enumerate(value):
                rowh = max(len(_wrap(_plain(cell), cellw - 12, size)) for cell in row) * size * 1.5 + 16
                for ci, cell in enumerate(row):
                    x = pad + ci * cellw
                    s.rect(x, y, cellw, rowh, '#f6f8fa' if ri == 0 else '#ffffff', '#d0d7de')
                    s.paragraph(x + 6, y + 8, _plain(cell), size, '#24292f', cellw - 12)
                y += rowh
        elif kind == 'rule':
            s.line(pad, y, w - pad, y, '#d0d7de'); y += size
        else:
            if kind in {'code', 'quote'}:
                used = len(_wrap(_plain(value), w - pad * 2 - 20, fs, 'mono' if kind == 'code' else 'sans')) * fs * 1.5 + 16
                s.rect(pad, y, w - pad * 2, used, '#f6f8fa')
                if kind == 'quote':
                    s.rect(pad, y, 4, used, '#d0d7de')
                s.paragraph(pad + 10, y + 8, _plain(value), fs, '#24292f', w - pad * 2 - 20, 'mono' if kind == 'code' else 'sans')
            else:
                used = s.paragraph(pad, y, ('• ' if kind == 'list' else '') + _plain(value), fs, '#24292f', w - pad * 2)
            y += used
        y += size
    s.height = max(120, math.ceil(y + pad))
    if s.width * s.height * p['scale'] ** 2 > 16_000_000:
        raise ValueError('文档长图超过 1600 万像素；请导出 PDF 或拆分文稿')
    return s


def _pdf(blocks, title):
    try:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted, HRFlowable
    except ImportError:
        raise RuntimeError('PDF 转换需要安装 reportlab') from None
    font_name = 'HouSignal-Unicode'
    if font_name not in pdfmetrics.getRegisteredFontNames():
        # Embed a real Unicode subset so every PDF viewer sees the same glyphs.
        # External CID-font substitution can turn bullets into unrelated CJK.
        for candidate in (
            '/System/Library/Fonts/Supplemental/Arial Unicode.ttf',
            '/Library/Fonts/Arial Unicode.ttf',
            '/usr/share/fonts/truetype/noto/NotoSansSC-Regular.ttf',
            '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
            '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',
            'C:/Windows/Fonts/msyh.ttc',
        ):
            if Path(candidate).is_file():
                try:
                    font = TTFont(font_name, candidate, subfontIndex=0)
                    if ord('中') not in font.face.charToGlyph or ord('•') not in font.face.charToGlyph:
                        continue
                    pdfmetrics.registerFont(font)
                    break
                except Exception:
                    continue
        else:
            raise ValueError('便携PDF需本项目安装支持中文的可嵌入TrueType字体')
    styles = getSampleStyleSheet()
    base = ParagraphStyle('NativeBody', fontName=font_name, fontSize=11, leading=18, alignment=TA_LEFT, spaceAfter=10)
    story = []
    for kind, value, meta in blocks:
        if kind == 'table':
            table = Table([[Paragraph(escape(cell), base) for cell in row] for row in value], repeatRows=1)
            table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#d0d7de')), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f6f8fa')), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story.extend([table, Spacer(1, 12)])
        elif kind == 'code':
            story.append(Preformatted(value, ParagraphStyle('NativeCode', parent=base, fontName=font_name, fontSize=9, leading=13), maxLineLength=85))
        elif kind == 'rule':
            story.append(HRFlowable(width='100%', color=colors.HexColor('#d0d7de')))
        else:
            style = ParagraphStyle('NativeHeading', parent=base, fontSize=20 if meta == 1 else 15, leading=26, spaceBefore=14) if kind == 'heading' else base
            story.append(Paragraph(escape(('• ' if kind == 'list' else '') + _plain(value)).replace('\n', '<br/>'), style))
    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title, author='HouSignal', topMargin=45, bottomMargin=45, leftMargin=40, rightMargin=40).build(story)
    return buf.getvalue()


def _docx(blocks):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError('DOCX 转换需要 python-docx') from None
    doc = Document()
    for kind, value, meta in blocks:
        if kind == 'heading':
            doc.add_heading(_plain(value), level=min(meta, 9))
        elif kind == 'table':
            table = doc.add_table(rows=len(value), cols=len(value[0]))
            table.style = 'Table Grid'
            for ri, row in enumerate(value):
                for ci, cell in enumerate(row):
                    table.cell(ri, ci).text = _plain(cell)
        else:
            doc.add_paragraph(_plain(value), style='List Bullet' if kind == 'list' else None)
    buf = BytesIO(); doc.save(buf)
    return buf.getvalue()


def _report(rows, p, out):
    columns = list(rows[0])
    numeric = [c for c in columns if any(isinstance(row[c], (int, float)) for row in rows) and all(row[c] is None or isinstance(row[c], (int, float)) for row in rows)]
    stats = {}
    for col in numeric:
        vals = [row[col] for row in rows if row[col] is not None]
        stats[col] = {'min': min(vals), 'max': max(vals), 'sum': sum(vals), 'mean': statistics.mean(vals),
            'median': statistics.median(vals), 'std': statistics.stdev(vals) if len(vals) > 1 else 0}
    categorical = [c for c in columns if c not in numeric]
    overview = {'shape': {'rows': len(rows), 'cols': len(columns)}, 'numeric_columns': numeric,
        'categorical_columns': categorical, 'numeric_stats': stats,
        'columns': [{'name': c, 'type': 'numeric' if c in numeric else 'categorical',
            'missing': sum(row[c] is None for row in rows), 'unique': len({row[c] for row in rows})} for c in columns],
        'total_missing': sum(row[c] is None for row in rows for c in columns),
        'top_categories': {c: [{'value': key, 'count': count} for key, count in Counter(str(row[c]) for row in rows if row[c] is not None).most_common(5)] for c in categorical}}
    kpi = p.get('kpi', numeric[:4])
    if not isinstance(kpi, list) or not all(c in stats for c in kpi):
        raise ValueError('kpi 必须是有效数值列数组')
    out.json('report.json', overview)
    if p['op'] == 'analyze' or p.get('format') == 'json':
        return
    fragments = ['<section style="padding:24px;background:white">', f'<h1>{escape(p.get("title", "数据可视化报告"))}</h1>',
        f'<p>{len(rows)} 行 × {len(columns)} 列 · 缺失值 {overview["total_missing"]} 个</p>', '<div style="display:flex;gap:24px">']
    for col in kpi:
        fragments.append(f'<section><h2>{stats[col]["sum"]:,.2f}</h2><p>{escape(col)} · 均值 {stats[col]["mean"]:,.2f}</p></section>')
    fragments.append('</div>')
    if numeric:
        value = numeric[0]
        if categorical:
            category = categorical[0]
            grouped = {}
            for row in rows:
                if row[value] is not None:
                    label = str(row[category] if row[category] is not None else '空值')
                    grouped[label] = grouped.get(label, 0) + row[value]
            chart_rows = [{'category': key, 'value': val} for key, val in sorted(grouped.items(), key=lambda item: item[1], reverse=True)[:10]]
        else:
            chart_rows = [{'category': col, 'value': stats[col]['sum']} for col in numeric[:8]]
        if chart_rows:
            scene = _chart(chart_rows, {**p, 'chart_type': 'column', 'title': value})
            _save_scene(out, 'report_chart', scene, {**p, 'format': 'all'}, html=False)
            fragments.append(scene.svg())
            if 2 <= len(chart_rows) <= 8 and all(row['value'] >= 0 for row in chart_rows) and sum(row['value'] for row in chart_rows) > 0:
                pie = _chart(chart_rows, {**p, 'chart_type': 'pie', 'title': value + ' 占比'})
                _save_scene(out, 'report_pie', pie, {**p, 'format': 'all'}, html=False)
                fragments.append(pie.svg())
    fragments.append('<table><thead><tr>' + ''.join(f'<th>{escape(col)}</th>' for col in columns) + '</tr></thead><tbody>')
    for row in rows[:20]:
        fragments.append('<tr>' + ''.join(f'<td>{escape(str(row[c]) if row[c] is not None else "")}</td>' for c in columns) + '</tr>')
    fragments.append('</tbody></table></section>')
    out.write('report.html', _html(p.get('title', '数据可视化报告'), ''.join(fragments)))


def _infographic(text, p, paths, out):
    op = p['op']
    data = _load_data(text, p, paths) if p.get('data') is not None or op in {'bar-race', 'count-up', 'progress', 'line-grow'} else None
    if op == 'hierarchy':
        tree = data or _outline(text, p.get('title'))
        scene = _tree_scene(tree, p)
        _save_scene(out, 'infographic', scene, p)
        out.json('infographic.json', tree)
        return
    if op in {'bar-race', 'count-up', 'progress', 'line-grow'}:
        _animation(data, p, out)
        return
    if data is not None:
        if isinstance(data, Mapping):
            data = data.get('items', data.get('data'))
        if not isinstance(data, list) or not data or len(data) > 30:
            raise ValueError('静态信息图 data 需要 1–30 个 items')
        items = []
        for item in data:
            if isinstance(item, Mapping):
                label = item.get('label', item.get('name', ''))
                description = item.get('description', item.get('desc', ''))
                label = _string(label, 'item.label', maximum=300, empty=False)
                description = _string(description, 'item.description', maximum=1000)
                items.append(label + (' — ' + description if description else ''))
            else:
                items.append(_string(item, 'item', maximum=1000, empty=False))
    else:
        lines = [line.strip().lstrip('#-* ').strip() for line in text.splitlines() if line.strip()]
        if not lines:
            raise ValueError('信息图内容为空')
        items = lines[1:] if text.lstrip().startswith('#') else lines
        if not items:
            items = lines
    w, h, pad = p['width'], p['height'], 30
    s = _Scene(w, h)
    size = max(12, round(w / 44))
    title = p.get('title') or (text.splitlines()[0].lstrip('# ').strip() if text else '')
    if title:
        s.text(pad, pad, title, size * 1.6, '#24292f')
    top = pad + size * 3
    count = len(items)
    horizontal = op == 'flow' and count <= 6
    slotw = (w - pad * 2) / count if horizontal else (w - pad * 2) / (2 if op == 'comparison' else 1)
    sloth = h - top - pad if horizontal else (h - top - pad) / math.ceil(count / (2 if op == 'comparison' else 1))
    for i, item in enumerate(items):
        x = pad + i * slotw if horizontal else pad + (i % 2) * slotw if op == 'comparison' else pad
        y = top if horizontal else top + (i // 2 if op == 'comparison' else i) * sloth
        width = slotw - 16
        lines = _wrap(item, width - size * 2, size)
        if len(lines) * size * 1.5 + size * 4 > sloth:
            raise ValueError('信息图文字超出画幅；请增加尺寸或拆分条目')
        s.rect(x, y, width, sloth - 16, '#f6f8fa', '#d0d7de')
        s.text(x + size, y + size, f'{i + 1:02d}', size * 1.5, _COLORS[i % len(_COLORS)])
        s.paragraph(x + size, y + size * 3, item, size, '#24292f', width - size * 2)
        if op == 'flow' and i < count - 1:
            if horizontal:
                ax, ay = x + width, y + sloth / 2
                s.polygon([(ax + 3, ay - 6), (ax + 14, ay), (ax + 3, ay + 6)], _COLORS[i % len(_COLORS)])
            else:
                ax, ay = x + width / 2, y + sloth - 14
                s.polygon([(ax - 6, ay), (ax + 6, ay), (ax, ay + 10)], _COLORS[i % len(_COLORS)])
    _save_scene(out, 'infographic', s, p)
    out.json('infographic.json', {'op': op, 'items': items})


def _animation(data, p, out):
    if not isinstance(data, Mapping):
        raise ValueError('动画数据需要 JSON 对象')
    fps = _number(p.get('fps', 20), 'fps', 1, 30, True)
    duration = _number(p.get('duration', 4), 'duration', .2, 15)
    count = max(2, round(fps * duration))
    if count > 150 or p['width'] * p['height'] * p['scale'] ** 2 * count > 100_000_000:
        raise ValueError('动画超过 150 帧或总像素预算；降低画幅、scale、fps 或 duration')
    op = p['op']
    title = p.get('title') or _string(data.get('title', data.get('label', '')), 'title', maximum=500)
    width, height = p['width'], p['height']
    top_n = _number(p.get('top_n', 10), 'top_n', 1, 30, True)
    if op in {'bar-race', 'line-grow'}:
        xs = data.get('times' if op == 'bar-race' else 'x')
        series = data.get('series')
        if not isinstance(xs, list) or not 2 <= len(xs) <= 100 or not isinstance(series, Mapping) or not 1 <= len(series) <= 30:
            raise ValueError('动画需要至少两时间点和非空 series')
        xs = [_string(str(x), 'time', maximum=100, empty=False) for x in xs]
        normalized = {}
        for name, values in series.items():
            _string(name, 'series 名称', maximum=100, empty=False)
            if not isinstance(values, list) or len(values) != len(xs):
                raise ValueError('series 长度必须等于时间点数量')
            normalized[name] = [_number(value, 'series value') for value in values]
        series = normalized
    elif op == 'count-up':
        items = data.get('items', [data])
        if not isinstance(items, list) or not 1 <= len(items) <= 8:
            raise ValueError('count-up 需要 1–8 个 items')
        items = [{'label': _string(item.get('label', ''), 'label', maximum=100),
            'value': _number(item.get('value'), 'value'),
            'prefix': _string(item.get('prefix', ''), 'prefix', maximum=20),
            'suffix': _string(item.get('suffix', ''), 'suffix', maximum=20)} for item in items if isinstance(item, Mapping)]
        if len(items) != len(data.get('items', [data])):
            raise ValueError('count-up item 必须是对象')
    else:
        value = _number(data.get('value'), 'value', 0)
        maximum = _number(data.get('max', 100), 'max', .00001)
        if value > maximum:
            raise ValueError('progress value 不能大于 max')
        color = _color(data.get('color', _COLORS[0]))
        style = p.get('style', 'ring')
        if style not in {'ring', 'bar'}:
            raise ValueError('progress style 必须是 ring/bar')
    out.check(['chart.gif', 'chart.json'])
    frames = []
    for index in range(count):
        t = index / (count - 1)
        eased = 1 - (1 - t) ** 3
        s = _Scene(width, height, '#0f1420' if op == 'count-up' else '#ffffff')
        size = max(12, round(width / 30))
        if title:
            s.text(20, 12, title, size, '#8a94a6' if op == 'count-up' else '#24292f')
        if op == 'count-up':
            for i, item in enumerate(items):
                yy = 45 + i * (height - 55) / len(items)
                target = item['value']
                digits = 0 if float(target).is_integer() else 1
                s.text(20, yy, f'{item["prefix"]}{target * eased:,.{digits}f}{item["suffix"]}', size * 2, _COLORS[i % len(_COLORS)])
                s.text(width * .62, yy + size / 2, item['label'], size, '#c8cfdb')
        elif op == 'progress':
            fraction = value / maximum * eased
            if style == 'bar':
                s.rect(20, height * .45, width - 40, height * .18, '#e9edf2')
                s.rect(20, height * .45, (width - 40) * fraction, height * .18, color)
            else:
                cx, cy, radius = width / 2, height * .55, min(width, height) * .32
                s.ellipse(cx - radius, cy - radius, radius * 2, radius * 2, '#e9edf2')
                end = -math.pi / 2 + fraction * math.tau
                samples = max(2, round(fraction * 100))
                s.polygon([(cx, cy), *[(cx + radius * math.cos(-math.pi / 2 + (end + math.pi / 2) * j / samples), cy + radius * math.sin(-math.pi / 2 + (end + math.pi / 2) * j / samples)) for j in range(samples + 1)]], color)
                r = radius * .72
                s.ellipse(cx - r, cy - r, r * 2, r * 2, '#ffffff')
            s.text(width * .4, height * .45, f'{fraction:.0%}', size * 2, '#24292f')
        elif op == 'bar-race':
            progress = t * (len(xs) - 1)
            seg = min(int(progress), len(xs) - 2)
            local = 1 - (1 - (progress - seg)) ** 3
            current = [(name, values[seg] + (values[seg + 1] - values[seg]) * local) for name, values in series.items()]
            ranked = sorted(current, key=lambda item: item[1], reverse=True)[:top_n]
            s = _chart([{'category': name, 'value': value} for name, value in ranked], {**p, 'chart_type': 'bar', 'title': title + ' · ' + xs[min(round(progress), len(xs) - 1)]})
        else:
            progress = t * (len(xs) - 1)
            part = int(progress)
            low = min(0, min(min(vals) for vals in series.values()))
            high = max(max(vals) for vals in series.values())
            left, right, top, bottom = 45, width - 20, 45, height - 35
            s.line(left, top, left, bottom, '#d0d7de'); s.line(left, bottom, right, bottom, '#d0d7de')
            for si, (name, vals) in enumerate(series.items()):
                points = [(left + j / (len(xs) - 1) * (right - left), bottom - (vals[j] - low) / (high - low or 1) * (bottom - top)) for j in range(part + 1)]
                if part < len(xs) - 1 and progress > part:
                    val = vals[part] + (vals[part + 1] - vals[part]) * (progress - part)
                    points.append((left + progress / (len(xs) - 1) * (right - left), bottom - (val - low) / (high - low or 1) * (bottom - top)))
                for a, b in zip(points, points[1:]):
                    s.line(*a, *b, _COLORS[si % len(_COLORS)], 3)
                s.text(left + si * width / len(series), height - size - 2, name, size, _COLORS[si % len(_COLORS)])
        frames.append(s.image(p['scale']).convert('P', palette=0, colors=128))
    buffer = BytesIO()
    frames[0].save(buffer, format='GIF', save_all=True, append_images=frames[1:],
        duration=[round(1000 / fps)] * (len(frames) - 1) + [800], loop=0, optimize=True, disposal=2)
    out.write('chart.gif', buffer.getvalue())
    out.json('chart.json', {'op': op, 'fps': fps, 'duration': duration, 'data': data})


def dispatch(skill_name, *, input_text, assets, parameters, output_dir, context):
    """Execute a registered local operation and return real, scoped artifacts.

    assets must already be authorized and resolved to Paths by the project API.
    context is accepted for the dispatcher contract, never used to fetch secrets,
    global profiles, credentials or account cookies.
    """
    if skill_name not in SKILL_NAMES:
        raise ValueError('未注册此文档技能')
    p = _validate(skill_name, parameters)
    paths = _asset_paths(assets)
    text = _string(input_text, 'input_text', maximum=100000)
    # Validate source content before creating the caller's output directory.
    if skill_name == 'card-design' and p['op'] == 'audit':
        report = _audit(paths)
        out = _Output(output_dir); out.json('card_audit.json', report)
        return {'response': f'已检查 {len(report)} 张卡片；{sum(r["pass"] for r in report)} 张通过密度检查。', 'artifact_paths': out.paths}
    if skill_name in {'comparison-card', 'chart-visualization', 'data-report'}:
        data = _load_data(text, p, paths)
        # Chart validation is run before any artifact is written.
        if skill_name == 'chart-visualization':
            scene, rows = _prepare_chart(data, p)
        else:
            rows = _rows(data)
    elif skill_name != 'infographic':
        text = _input(text, paths)
        if not text and not (skill_name in CARD_SIZES and p.get('cards')) and not (skill_name == 'mindmap' and p.get('data')):
            raise ValueError('内容不能为空')
    out = _Output(output_dir)
    if skill_name in CARD_SIZES and skill_name != 'comparison-card':
        _cards(text, p, out)
    elif skill_name == 'comparison-card':
        _comparison(rows, p, out)
    elif skill_name == 'chart-visualization':
        out.check(['chart.svg', 'chart.png', 'chart.html', 'chart.json'])
        _save_scene(out, 'chart', scene, p)
        out.json('chart.json', {'chart_type': p['chart_type'], 'data': rows})
    elif skill_name == 'data-report':
        _report(rows, p, out)
    elif skill_name == 'mindmap':
        tree = p.get('data') or _outline(text, p.get('title'))
        scene = _tree_scene(tree, p)
        out.check(['mindmap.svg', 'mindmap.png', 'mindmap.html', 'mindmap.json'])
        _save_scene(out, 'mindmap', scene, p, interactive=True)
        out.json('mindmap.json', tree)
    elif skill_name == 'infographic':
        _infographic(text, p, paths, out)
    else:
        blocks = _blocks(text)
        title = p.get('title') or text.splitlines()[0].lstrip('# ').strip()
        fmt = p.get('format', 'html')
        prefix = 'article' if skill_name == 'gzh-design' else 'document'
        out.check([prefix + '.html', prefix + '.' + fmt])
        body = _document_html(blocks, title, p, gzh=skill_name == 'gzh-design')
        if fmt == 'pdf':
            converted = _pdf(blocks, title)
        elif fmt == 'docx':
            converted = _docx(blocks)
        elif fmt == 'png':
            scene = _document_scene(blocks, p)
            buf = BytesIO(); scene.image(p['scale']).save(buf, format='PNG'); converted = buf.getvalue()
        out.write(prefix + '.html', body)
        if fmt != 'html':
            out.write(prefix + '.' + fmt, converted)
    return {'response': f'已在楼响本地生成 {len(out.paths)} 份真实产物；请核对内容后使用。', 'artifact_paths': out.paths}
